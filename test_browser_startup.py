"""Actual Flask '/' and Chrome/CDP regressions, never the configured real database.

Run: python -m pytest test_browser_startup.py -v
Chrome: BOM_TEST_CHROME (default installed Windows Chrome absolute path).
Node: BOM_TEST_NODE (default node; native WebSocket required).
CDN: dependencies are fetched once into pytest's temporary directory, then served
by CDP at their EXACT original URLs. BOM_TEST_CDN_CACHE may point to a prefilled
cache (filenames below) for offline runs. App HTML/JS/API success are NEVER mocked.
Missing Chrome/Node/CDN dependencies FAIL explicitly, rather than silently skip.
Artifacts: per-test screenshot, DOM and CDP diagnostics in pytest temporary dirs.
"""
import base64
from contextlib import closing, contextmanager
import gc
import json
from io import BytesIO
import os
from pathlib import Path
import queue
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
from urllib.request import Request, urlopen

import pytest

ROOT = Path(__file__).resolve().parent
CHROME = Path(os.environ.get('BOM_TEST_CHROME', 'C:/Program Files/Google/Chrome/Application/chrome.exe'))
DEPENDENCIES = {
    'https://cdnjs.cloudflare.com/ajax/libs/jquery/3.6.0/jquery.min.js': ('jquery-3.6.0.min.js', 'application/javascript'),
    'https://cdnjs.cloudflare.com/ajax/libs/jstree/3.3.12/jstree.min.js': ('jstree-3.3.12.min.js', 'application/javascript'),
    'https://cdnjs.cloudflare.com/ajax/libs/jstree/3.3.12/themes/default/style.min.css': ('jstree-3.3.12-style.min.css', 'text/css'),
    'https://cdnjs.cloudflare.com/ajax/libs/jstree/3.3.12/themes/default/32px.png': ('jstree-3.3.12-32px.png', 'image/png'),
    'https://cdnjs.cloudflare.com/ajax/libs/jstree/3.3.12/themes/default/40px.png': ('jstree-3.3.12-40px.png', 'image/png'),
    'https://cdnjs.cloudflare.com/ajax/libs/jstree/3.3.12/themes/default/throbber.gif': ('jstree-3.3.12-throbber.gif', 'image/gif'),
}


def read_url(url, timeout=10):
    with urlopen(Request(url, headers={'User-Agent': 'BOM startup regression'}), timeout=timeout) as response:
        return response.read()


def stop_process(process):
    """Terminate only this owned process; never kill by image name or port."""
    if process.poll() is None:
        process.terminate()
        try:
            process.wait(timeout=8)
        except subprocess.TimeoutExpired:
            process.kill()
            process.wait(timeout=8)


FIXTURE_DATABASE_DIRS = []


@pytest.fixture(scope='session', autouse=True)
def cleanup_fixture_databases():
    yield
    gc.collect()
    for directory in FIXTURE_DATABASE_DIRS:
        shutil.rmtree(directory)


@contextmanager
def isolated_server(directory, database=None, legacy=False, anonymous_auth=False):
    directory.mkdir(parents=True, exist_ok=True)
    # Retain the isolated DB until suite teardown for two-process restart tests.
    # Workspace --basetemp contains browser artifacts, never fixture databases.
    if database is None:
        external = Path(tempfile.mkdtemp(prefix='bom-browser-db-'))
        FIXTURE_DATABASE_DIRS.append(external)
        database = external / 'fixture.sqlite'
    ready = directory / 'server-ready.json'
    command = [sys.executable, str(ROOT / 'tests' / '_browser_startup_server.py'),
               '--database', str(database), '--ready', str(ready)]
    if anonymous_auth:
        # Scope only the harness's automatic engineer identity: real auth routes,
        # templates, storage, HTTP and scripts remain untouched. Main '/' still
        # uses the real engineer fixture; auth GETs represent a logged-out user.
        bootstrap = (
            'import runpy, sys; from flask_login import LoginManager; '
            'original = LoginManager.request_loader; '
            'LoginManager.request_loader = lambda manager, callback: original(manager, '
            'lambda request: None if request.path == "/login" or request.path.startswith("/auth/") '
            'else callback(request)); '
            'sys.argv = sys.argv[1:]; runpy.run_path(sys.argv[0], run_name="__main__")'
        )
        command = [sys.executable, '-c', bootstrap, *command[1:]]
    if legacy:
        command.append('--legacy')
    with (directory / 'server.log').open('w', encoding='utf-8') as log:
        process = subprocess.Popen(command, cwd=str(ROOT), stdout=log, stderr=subprocess.STDOUT)
        try:
            deadline = time.monotonic() + 40
            while not ready.exists():
                if process.poll() is not None or time.monotonic() > deadline:
                    log.flush()
                    raise AssertionError('Isolated Flask startup failed: ' + (directory / 'server.log').read_text(encoding='utf-8'))
                time.sleep(.05)
            info = json.loads(ready.read_text(encoding='utf-8'))
            assert Path(info['database']).resolve() == database.resolve()
            yield info['url'], database
        finally:
            stop_process(process)


@pytest.fixture(scope='session')
def cdn_assets(tmp_path_factory):
    directory = tmp_path_factory.mktemp('startup-cdn')
    configured = os.environ.get('BOM_TEST_CDN_CACHE')
    result = {}
    for url, (filename, mime) in DEPENDENCIES.items():
        path = Path(configured) / filename if configured else directory / filename
        if not path.exists():
            if configured:
                pytest.fail(f'Offline CDN cache is incomplete: {path}. Startup tests were NOT skipped.')
            try:
                path.write_bytes(read_url(url, timeout=40))
            except Exception as error:
                pytest.fail(f'Cannot obtain real CDN dependency {url}: {error}. Set BOM_TEST_CDN_CACHE to a prefilled cache. Startup tests were NOT skipped.')
        body = path.read_bytes()
        assert len(body) > 20, f'Invalid CDN asset: {path}'
        if mime == 'application/javascript':
            assert not body.lstrip().startswith(b'<'), f'CDN returned HTML instead of JavaScript: {path}'
        result[url] = (body, mime)
    return result


class Browser:
    def __init__(self, directory, assets):
        self.directory = directory
        self.assets = assets
        self.events = []
        self.incoming = queue.Queue()
        self.sequence = 0
        self.server = None
        self.chrome = None
        self.transport = None
        self.logs = []
        self.data_mode = 'normal'
        self.paused_data = []
        self.failed_data = 0
        self.data_requests = 0
        self.hold_node_save = False
        self.paused_node_saves = []
        self.paused_exports = []

    def __enter__(self):
        self.directory.mkdir(parents=True, exist_ok=True)
        try:
            if not CHROME.is_absolute() or not CHROME.is_file():
                raise AssertionError(f'Installed Chrome unavailable: {CHROME}; set BOM_TEST_CHROME. Tests are NOT skipped.')
            node = shutil.which(os.environ.get('BOM_TEST_NODE', 'node'))
            if not node:
                raise AssertionError('Installed Node unavailable; set BOM_TEST_NODE. Tests are NOT skipped.')
            check = subprocess.run([node, '-e', 'if(typeof WebSocket!=="function")process.exit(1)'], capture_output=True)
            assert check.returncode == 0, 'Node native WebSocket unavailable (Node 22+ required); tests are NOT skipped'
            chrome_log = (self.directory / 'chrome.log').open('w', encoding='utf-8')
            self.logs.append(chrome_log)
            profile = self.directory / 'chrome-profile'
            self.chrome = subprocess.Popen([str(CHROME), '--headless=new', '--disable-gpu',
                '--no-first-run', '--no-default-browser-check', '--remote-debugging-port=0',
                '--remote-debugging-address=127.0.0.1', '--user-data-dir=' + str(profile),
                '--window-size=1600,1000', 'about:blank'], stdout=chrome_log, stderr=subprocess.STDOUT)
            port_file = profile / 'DevToolsActivePort'
            deadline = time.monotonic() + 25
            port = None
            while port is None:
                assert self.chrome.poll() is None, 'Chrome exited before CDP became available'
                assert time.monotonic() < deadline, 'Chrome CDP port file timed out'
                # Chrome briefly holds an exclusive Windows handle while writing.
                # Existence is not readiness: wait for a readable complete port.
                try:
                    lines = port_file.read_text().splitlines()
                    if lines and lines[0].isdigit():
                        port = lines[0]
                except OSError:
                    pass
                if port is None:
                    threading.Event().wait(.05)
            targets = json.loads(read_url(f'http://127.0.0.1:{port}/json/list'))
            target = next(t for t in targets if t['type'] == 'page')
            transport_log = (self.directory / 'transport.log').open('w', encoding='utf-8')
            self.logs.append(transport_log)
            self.transport = subprocess.Popen([node, str(ROOT / 'tests' / '_browser_startup_cdp.js'), target['webSocketDebuggerUrl']],
                stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=transport_log, text=True, encoding='utf-8', bufsize=1)
            def reader():
                for line in self.transport.stdout:
                    try:
                        self.incoming.put(json.loads(line))
                    except ValueError:
                        self.incoming.put({'transportError': line})
                self.incoming.put({'transportError': 'CDP transport closed'})
            self.reader = threading.Thread(target=reader, daemon=True)
            self.reader.start()
            assert self.incoming.get(timeout=15).get('ready'), 'CDP transport failed to connect'
            for domain in ['Page', 'Runtime', 'Log', 'Network']:
                self.call(domain + '.enable')
            self.call('Network.setCacheDisabled', {'cacheDisabled': True})
            self.call('Fetch.enable', {'patterns': [
                {'urlPattern': 'https://cdnjs.cloudflare.com/*', 'requestStage': 'Request'},
                {'urlPattern': '*/api/data', 'requestStage': 'Request'},
                {'urlPattern': '*/api/node/*', 'requestStage': 'Request'},
                {'urlPattern': '*/export/excel*', 'requestStage': 'Response'}]})
            return self
        except BaseException:
            self.close()
            raise

    def send(self, method, params=None):
        self.sequence += 1
        self.transport.stdin.write(json.dumps({'id': self.sequence, 'method': method, 'params': params or {}}) + '\n')
        self.transport.stdin.flush()
        return self.sequence

    def dispatch(self, message):
        assert 'transportError' not in message, message
        if 'method' not in message:
            return
        self.events.append(message)
        if message['method'] != 'Fetch.requestPaused':
            return
        params = message['params']
        request_id = params['requestId']
        url = params['request']['url']
        if url in self.assets:
            body, mime = self.assets[url]
            self.send('Fetch.fulfillRequest', {'requestId': request_id, 'responseCode': 200,
                'responseHeaders': [{'name': 'Content-Type', 'value': mime}],
                'body': base64.b64encode(body).decode('ascii')})
        elif url.endswith('/api/data'):
            self.data_requests += 1
            if self.data_mode == 'hold':
                self.paused_data.append(request_id)
            elif self.data_mode == 'fail_once' and not self.failed_data:
                self.failed_data += 1
                self.send('Fetch.fulfillRequest', {'requestId': request_id, 'responseCode': 503,
                    'responseHeaders': [{'name': 'Content-Type', 'value': 'application/json'}],
                    'body': base64.b64encode(b'{"error":"startup test controlled outage"}').decode('ascii')})
            elif self.data_mode == 'malformed_once' and not self.failed_data:
                self.failed_data += 1
                self.send('Fetch.fulfillRequest', {'requestId': request_id, 'responseCode': 200,
                    'responseHeaders': [{'name': 'Content-Type', 'value': 'application/json'}],
                    'body': base64.b64encode(b'{"nodes":[],"root_ids":"invalid"}').decode('ascii')})
            else:
                self.send('Fetch.continueRequest', {'requestId': request_id})
        elif '/export/excel' in url:
            self.paused_exports.append(params)
        elif '/api/node/' in url:
            if self.hold_node_save and params['request']['method'] == 'PUT':
                self.paused_node_saves.append(params)
            else:
                self.send('Fetch.continueRequest', {'requestId': request_id})
        else:
            raise AssertionError(f'Uncached CDN resource {url}; declare the real dependency rather than mock it')

    def call(self, method, params=None, timeout=20):
        request_id = self.send(method, params)
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            try:
                message = self.incoming.get(timeout=max(.01, deadline - time.monotonic()))
            except queue.Empty:
                break
            self.dispatch(message)
            if message.get('id') == request_id:
                assert 'error' not in message, message
                return message.get('result', {})
        raise AssertionError(f'CDP timeout: {method}')

    def evaluate(self, expression):
        response = self.call('Runtime.evaluate', {'expression': expression, 'returnByValue': True, 'awaitPromise': True})
        assert 'exceptionDetails' not in response, response
        return response.get('result', {}).get('value')

    def wait(self, expression, timeout=20):
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            # DOM objects serialize as {}, which is false in Python despite being
            # truthy in JS. Convert in the browser, not after CDP serialization.
            if self.evaluate('Boolean(' + expression + ')'):
                return
            time.sleep(.08)
        raise AssertionError(f'Browser condition timed out: {expression}\nDOM: ' + str(self.evaluate('document.body.innerText'))[:3500])

    def navigate(self, url):
        self.call('Page.navigate', {'url': url})

    def release_data(self):
        self.data_mode = 'normal'
        for request_id in self.paused_data:
            self.call('Fetch.continueRequest', {'requestId': request_id})
        self.paused_data.clear()

    def screenshot(self, name):
        result = self.call('Page.captureScreenshot', {'format': 'png', 'captureBeyondViewport': False})
        path = self.directory / (name + '.png')
        path.write_bytes(base64.b64decode(result['data']))
        return path

    def assert_clean(self, expected_data_failure=False, script='script.js'):
        # Drain pending browser events via a roundtrip; errors remain evidence.
        self.evaluate('document.readyState')
        exceptions = [e for e in self.events if e['method'] == 'Runtime.exceptionThrown']
        console_errors = [e for e in self.events if e['method'] == 'Runtime.consoleAPICalled'
                          and e['params']['type'] in ('error', 'assert')]
        log_errors = [e for e in self.events if e['method'] == 'Log.entryAdded'
                      and e['params']['entry']['level'] == 'error'
                      and not (expected_data_failure and '/api/data' in e['params']['entry'].get('url', ''))
                      and '/favicon.ico' not in e['params']['entry'].get('url', '')]
        if expected_data_failure:
            console_errors = [e for e in console_errors if not any(
                re.search(r'503|load|بارگذاری', a.get('value', '') if isinstance(a.get('value'), str) else a.get('description', ''), re.I)
                for a in e['params']['args'])]
        assert not exceptions, json.dumps(exceptions, ensure_ascii=False)
        assert not console_errors, json.dumps(console_errors, ensure_ascii=False)
        assert not log_errors, json.dumps(log_errors, ensure_ascii=False)
        assert any(e['method'] == 'Network.responseReceived' and
                   '/static/' + script in e['params']['response']['url'] and
                   e['params']['response']['status'] == 200 for e in self.events), 'Actual app script was not loaded'

    def close(self):
        if self.transport and self.transport.poll() is None:
            try:
                self.send('Browser.close')
                if self.chrome:
                    self.chrome.wait(timeout=8)
            except (OSError, BrokenPipeError, subprocess.TimeoutExpired):
                pass
        if self.transport:
            stop_process(self.transport)
            if self.transport.stdin:
                self.transport.stdin.close()
            if hasattr(self, 'reader'):
                self.reader.join(timeout=2)
            if self.transport.stdout:
                self.transport.stdout.close()
        if self.chrome:
            stop_process(self.chrome)
        for log in self.logs:
            log.close()

    def __exit__(self, exc_type, exc, traceback):
        try:
            if self.transport and self.transport.poll() is None:
                self.screenshot('failure' if exc else 'final')
                (self.directory / 'dom.html').write_text(self.evaluate('document.documentElement.outerHTML'), encoding='utf-8')
        finally:
            (self.directory / 'cdp-events.json').write_text(json.dumps(self.events, ensure_ascii=False, indent=2), encoding='utf-8')
            self.close()


VISIBLE = "e => !!e && e.getClientRects().length > 0 && getComputedStyle(e).visibility !== 'hidden'"


STYLE_AUDIT = r"""(selector, textToken, backgroundToken, pseudo) => {
    const e = document.querySelector(selector);
    if (!e) return null;
    const s = getComputedStyle(e, pseudo || null);
    const probe = document.createElement('span'); document.body.append(probe);
    const token = name => { probe.style.color = `var(${name})`; return getComputedStyle(probe).color; };
    const rgba = value => { const n = value.match(/[\d.]+/g).map(Number); return [...n.slice(0,3), n[3] ?? 1]; };
    const over = (top, bottom) => top.slice(0,3).map((c,i)=>c*top[3]+bottom[i]*(1-top[3])).concat(1);
    const chain = []; for (let p=e; p; p=p.parentElement) chain.unshift(p);
    let bg = [255,255,255,1];
    for (const p of chain) bg = over(rgba(getComputedStyle(p).backgroundColor), bg);
    const fg = over(rgba(s.color), bg);
    const lum = rgb => rgb.slice(0,3).map(c=>{c/=255;return c<=.04045?c/12.92:((c+.055)/1.055)**2.4;}).reduce((a,c,i)=>a+c*[.2126,.7152,.0722][i],0);
    const a=lum(fg), b=lum(bg);
    const result = {color:s.color, background:s.backgroundColor, scheme:s.colorScheme,
        textMatches:!textToken || s.color===token(textToken),
        backgroundMatches:!backgroundToken || s.backgroundColor===token(backgroundToken),
        contrast:(Math.max(a,b)+.05)/(Math.min(a,b)+.05),
        disabled:e.disabled, readonly:e.readOnly,
        settled:chain.every(p=>p.getAnimations().every(a=>a.playState!=='running')),
        visible:!!(e.tagName==='OPTION' ? e.closest('select') : e).getClientRects().length};
    probe.remove(); return result;
}"""


def assert_style(browser, selector, text_token, background_token=None, pseudo=None, disabled=False, scheme_selector='html'):
    expression = f'({STYLE_AUDIT})({json.dumps(selector)},{json.dumps(text_token)},{json.dumps(background_token)},{json.dumps(pseudo)})'
    # Wait on the actual interpolated computed colors, not inline values or a
    # guessed delay: input/body transitions can outlive the theme attribute.
    browser.wait(f'(() => {{const s={expression};return s && s.textMatches && s.backgroundMatches && s.settled;}})()')
    style = browser.evaluate(expression)
    assert style['visible'], (selector, style)
    assert style['scheme'] == browser.evaluate(f'getComputedStyle(document.querySelector({json.dumps(scheme_selector)})).colorScheme'), (selector, style)
    if disabled:
        assert style['disabled'], (selector, style)
    else:
        assert style['contrast'] >= 4.5, f'WCAG text contrast < 4.5: {selector}: {style}'
    return style


def switch_theme(browser, theme, selector='#theme-toggle'):
    dom_click(browser, selector)
    browser.wait(f"document.documentElement.dataset.theme === '{theme}' && getComputedStyle(document.documentElement).colorScheme === '{theme}' && localStorage.getItem('theme') === '{theme}'")


def assert_products(browser, names=None):
    names = names or {'p41': 'Startup Product Alpha', 'p73': 'محصول ماندگار Beta'}
    browser.wait("document.querySelector('#p41_anchor') && document.querySelector('#p73_anchor')")
    for node_id, name in names.items():
        assert browser.evaluate(f"document.getElementById({json.dumps(node_id + '_anchor')}).textContent.includes({json.dumps(name)})")
        assert browser.evaluate(f"({VISIBLE})(document.getElementById({json.dumps(node_id + '_anchor')}))")
    assert browser.evaluate("Array.from(document.querySelectorAll('#tree-container .jstree-anchor')).every(a=>!/(?:Required|Shortage)\\s*:/i.test(a.textContent))"), 'Tree labels should show status through color, not required/shortage counts'
    roots = browser.evaluate("Array.from(document.querySelectorAll('#tree-container > ul > li')).map(e=>e.id)")
    assert set(roots) == set(names), f'Wrong root IDs or missing real products: {roots}'
    assert browser.data_requests >= 1, 'Initialization never requested real /api/data'


def select_inventory(browser, node_id, stock, coefficient, total, shortage):
    # Real widget open + DOM clicks; do not set currentData/currentNodeId or call showEditForm.
    if node_id != 'p41':
        browser.evaluate("$('#tree-container').jstree(true).open_all(); true")
    browser.wait(f"document.getElementById('{node_id}_anchor')")
    browser.evaluate(f"document.getElementById('{node_id}_anchor').click(); true")
    browser.wait(f"document.getElementById('field-name').value === document.getElementById('{node_id}_anchor').textContent.trim() || (typeof currentNodeId !== 'undefined' && currentNodeId === '{node_id}')")
    browser.evaluate("document.getElementById('tab-btn-inventory').click(); true")
    browser.wait(f"({VISIBLE})(document.getElementById('tab-inventory'))")
    fields = browser.evaluate("Object.fromEntries(['quantity','required_quantity','total_required'].map(k=>[k,document.getElementById('field-'+k).value]))")
    assert int(fields['quantity']) == stock
    if coefficient is not None:
        assert int(fields['required_quantity']) == coefficient
        assert browser.evaluate(f"({VISIBLE})(document.getElementById('field-required_quantity'))")
    assert int(fields['total_required']) == total
    assert browser.evaluate(f"({VISIBLE})(document.getElementById('field-total_required'))")
    text = browser.evaluate("document.getElementById('shortage-text').innerText")
    assert str(shortage) in text, f'Shortage {shortage} not displayed: {text}'
    assert browser.evaluate(f"({VISIBLE})(document.getElementById('shortage-info'))")


def test_real_startup_products_and_inventory(tmp_path, cdn_assets):
    with isolated_server(tmp_path / 'server') as (url, database), Browser(tmp_path / 'browser', cdn_assets) as browser:
        actual = json.loads(read_url(url + '/api/data'))
        assert actual['root_ids'] == ['p41', 'p73']
        assert actual['nodes']['p41']['name'] == 'Startup Product Alpha'
        browser.navigate(url + '/')
        assert_products(browser)
        select_inventory(browser, 'p41', 2, None, 10, 8)
        select_inventory(browser, 'a51', 4, 3, 24, 20)
        select_inventory(browser, 'a52', 5, 2, 40, 35)
        select_inventory(browser, 'r61', 7, 4, 140, 133)
        for theme in ['dark', 'light']:
            switch_theme(browser, theme)
            assert_style(browser, '#field-quantity', '--control-text', '--control-bg')
            readonly = assert_style(browser, '#field-total_required', '--control-readonly-text', '--control-readonly-bg')
            assert readonly['readonly']
            assert_style(browser, '#shortage-info', '--feedback-danger-text', '--feedback-danger-bg')
            assert_style(browser, '#shortage-text', '--feedback-danger-text')
            assert_style(browser, '#field-supplier-email', '--placeholder-color', pseudo='::placeholder')
            # Exercise a disabled state on the existing visible editable control;
            # no fabricated DOM, data mutation, or autosave event is involved.
            browser.evaluate("document.getElementById('field-quantity').disabled=true; true")
            try:
                assert_style(browser, '#field-quantity', '--control-disabled-text', '--control-disabled-bg', disabled=True)
            finally:
                browser.evaluate("document.getElementById('field-quantity').disabled=false; true")
            browser.screenshot('selected-part-inventory-' + theme)
        browser.assert_clean()


def test_dark_tree_status_search_hover_and_kanban_borders(tmp_path, cdn_assets):
    with isolated_server(tmp_path / 'tree-server') as (url, _), Browser(tmp_path / 'tree-browser', cdn_assets) as browser:
        browser.navigate(url + '/')
        assert_products(browser)
        select_inventory(browser, 'r61', 7, 4, 140, 133)
        # Actual DOM autosave creates a missing-stock node alongside partial and
        # sufficient fixtures, without replacing the successful backend payload.
        dom_fill(browser, '#field-quantity', 0)
        api_until(browser, url + '/api/data', lambda d: d['nodes']['r61']['quantity'] == 0)
        switch_theme(browser, 'dark')
        names = {'p41': 'Startup Product Alpha', 'a51': 'Startup Assembly',
                 'a52': 'Startup Nested', 'r61': 'Startup Part', 'p73': 'محصول ماندگار Beta'}
        tokens = {node: '--status-warning-text' for node in names}
        tokens.update(r61='--status-danger-text', p73='--status-success-text')
        contrast_failures = []
        for language, order in [('fa', 'سفارش'), ('en', 'Order')]:
            if language == 'en':
                dom_click(browser, '#lang-toggle')
                browser.wait("document.documentElement.dir === 'ltr'")
            browser.evaluate("$('#tree-container').jstree(true).open_all(); true")
            for node, name in names.items():
                browser.wait(f"document.getElementById('{node}_anchor')")
                label = browser.evaluate(f"document.getElementById('{node}_anchor').textContent.trim()")
                # Known names, type icon, and product order are the only allowed
                # label content. This catches Persian and English need suffixes.
                remainder = label.replace(name, '').strip()
                if node == 'p41':
                    assert f'(10 {order})' in remainder, label
                    remainder = remainder.replace(f'(10 {order})', '').strip()
                assert not re.search(r'[\w\d]', remainder), f'Unexpected tree label suffix: {label}'
                assert_style(browser, f'#{node}_anchor', tokens[node])
            dom_fill(browser, '#tree-search', 'Startup Part')
            browser.evaluate("document.getElementById('tree-search').dispatchEvent(new KeyboardEvent('keyup',{key:'t',bubbles:true})); true")
            browser.wait("document.getElementById('r61_anchor').classList.contains('jstree-search')")
            dom_click(browser, '#p73_anchor')
            try:
                assert_style(browser, '#r61_anchor', tokens['r61'], '--surface-search')
            except AssertionError as error:
                if 'WCAG text contrast < 4.5' not in str(error):
                    raise
                contrast_failures.append(language + ': ' + str(error))
            browser.screenshot('dark-tree-search-' + language)
            dom_click(browser, '#r61_anchor')
            assert_style(browser, '#r61_anchor', tokens['r61'], '--surface-selected')
            # The app attaches hover handlers after rebuilding jsTree. Wait for
            # the actual listener, not its implementation's guessed timer delay.
            browser.wait("Boolean($._data(document.getElementById('r61_anchor'),'events')?.mouseover)")
            # CDP mouse movement exercises actual :hover + delayed hover card.
            bounds = browser.evaluate("(() => {const r=document.getElementById('r61_anchor').getBoundingClientRect();return {x:r.x+r.width/2,y:r.y+r.height/2};})()")
            browser.call('Input.dispatchMouseEvent', {'type': 'mouseMoved', 'x': 1, 'y': 1})
            browser.call('Input.dispatchMouseEvent', {'type': 'mouseMoved', **bounds})
            browser.wait(f"({VISIBLE})(document.getElementById('hover-card'))")
            assert_style(browser, '#hover-card-status', '--status-danger-text')
            assert_style(browser, '#r61_anchor', '--status-danger-text', '--surface-selected')
            browser.call('Input.dispatchMouseEvent', {'type': 'mouseMoved', 'x': 1, 'y': 1})
            dom_fill(browser, '#tree-search', '')
            browser.evaluate("document.getElementById('tree-search').dispatchEvent(new KeyboardEvent('keyup',{bubbles:true})); true")
            dom_click(browser, '#btn-view-kanban')
            for node, status, border in [('r61', 'not-started', 'rgb(244, 67, 54)'),
                                         ('a51', 'in-progress', 'rgb(255, 152, 0)')]:
                selector = f'.kanban-card[onclick="goToNode(\'{node}\')"]'
                browser.wait(f"document.querySelector({json.dumps(selector)})")
                actual = browser.evaluate(f"(() => {{const e=document.querySelector({json.dumps(selector)}),s=getComputedStyle(e);return {{status:e.classList.contains('status-{status}'),border:s.{'borderLeftColor' if language == 'en' else 'borderRightColor'},width:s.{'borderLeftWidth' if language == 'en' else 'borderRightWidth'}}};}})()")
                assert actual == {'status': True, 'border': border, 'width': '4px'}, actual
                assert_style(browser, selector + ' .kanban-card-name', '--text-primary')
            browser.screenshot('dark-kanban-' + language)
            dom_click(browser, '#btn-view-tree')
            dom_click(browser, '#r61_anchor')
        browser.assert_clean()
        assert not contrast_failures, '\n'.join(contrast_failures)


@pytest.mark.parametrize('route', ['/login', '/auth/register', '/auth/login'])
def test_real_auth_routes_restore_main_theme_and_control_contrast(tmp_path, cdn_assets, route):
    failures = []
    with isolated_server(tmp_path / 'auth-server', anonymous_auth=True) as (url, database), Browser(tmp_path / 'auth-browser', cdn_assets) as browser:
        for theme in ['dark', 'light']:
            browser.navigate(url + '/')
            assert_products(browser)
            # Persist preference only through the actual main-page control.
            switch_theme(browser, theme)
            select_inventory(browser, 'r61', 7, 4, 140, 133)
            assert_style(browser, '#field-quantity', '--control-text', '--control-bg')
            browser.screenshot('main-inventory-' + theme)
            requests = browser.data_requests
            browser.navigate(url + route)
            browser.wait(f"location.pathname === {json.dumps(route)} && document.readyState === 'complete' && document.querySelector('.auth-form #username')")
            assert browser.evaluate("!document.querySelector('.app-layout')"), 'Auth fixture must be genuinely logged out'
            state = browser.evaluate("({theme:document.documentElement.dataset.theme ?? null, saved:localStorage.getItem('theme'), scheme:getComputedStyle(document.body).colorScheme})")
            if state != {'theme': theme, 'saved': theme, 'scheme': theme}:
                failures.append(f'{route} {theme}: theme restoration failed: {state}')
            dom_fill(browser, '#username', 'Contrast fixture')
            selectors = browser.evaluate("Array.from(document.querySelectorAll('.auth-form input:not([type=checkbox])'),e=>'#'+e.id)")
            audits = []
            for selector in selectors:
                audits.extend([(selector, '--text', '--bg-card', None),
                               (selector, '--auth-placeholder', None, '::placeholder')])
            audits.extend([('.auth-form label', '--text-secondary', None, None),
                           ('.auth-header h1', '--text', None, None),
                           ('.auth-header p', '--text-muted', None, None),
                           ('.auth-footer', '--text-secondary', None, None),
                           ('.auth-footer a', '--auth-link', None, None),
                           ('.auth-form button[type=submit]', None, None, None)])
            for selector, text, background, pseudo in audits:
                try:
                    assert_style(browser, selector, text, background, pseudo, scheme_selector='body')
                except AssertionError as error:
                    failures.append(f'{route} {theme}: {error}')
            # These routes contain no select/option controls; do not fabricate
            # them and misrepresent that as real-route coverage.
            assert browser.evaluate("document.querySelectorAll('.auth-form select').length") == 0
            browser.screenshot('auth-' + route.strip('/').replace('/', '-') + '-' + theme)
            assert browser.data_requests == requests, 'Auth unexpectedly initialized BOM data'
            browser.assert_clean(script='modern/app.js')
        with closing(sqlite3.connect(database)) as connection:
            assert connection.execute('SELECT count(*) FROM users').fetchone()[0] == 1
        assert not failures, '\n'.join(failures)


def test_startup_displays_loading_until_real_data_arrives(tmp_path, cdn_assets):
    with isolated_server(tmp_path / 'loading') as (url, _), Browser(tmp_path / 'loading-browser', cdn_assets) as browser:
        browser.data_mode = 'hold'
        browser.navigate(url + '/')
        browser.wait('document.readyState === "complete"')
        assert browser.paused_data, 'Startup did not request /api/data'
        loading = browser.evaluate("document.querySelector('.container').innerText")
        assert re.search(r'loading|بارگذاری', loading, re.I), f'No visible startup loading state: {loading}'
        assert 'هیچ محصولی تعریف نشده' not in loading and 'No products defined' not in loading
        browser.screenshot('visible-loading')
        browser.release_data()
        assert_products(browser)
        browser.assert_clean()


def test_products_persist_across_two_server_processes(tmp_path, cdn_assets):
    with isolated_server(tmp_path / 'first') as (url, database), Browser(tmp_path / 'first-browser', cdn_assets) as browser:
        browser.navigate(url + '/')
        assert_products(browser)
        # Perform an actual DOM edit and verify the persisted HTTP response, not a mock fetch.
        browser.evaluate("document.getElementById('p41_anchor').click(); true")
        browser.wait("document.getElementById('field-name').value === 'Startup Product Alpha'")
        browser.evaluate("const field=document.getElementById('field-name'); field.value='Persisted DOM Product'; field.dispatchEvent(new Event('input',{bubbles:true})); field.dispatchEvent(new Event('change',{bubbles:true})); true")
        deadline = time.monotonic() + 15
        while json.loads(read_url(url + '/api/data'))['nodes']['p41']['name'] != 'Persisted DOM Product':
            assert time.monotonic() < deadline, 'Real DOM edit did not persist through actual autosave API'
            browser.evaluate('document.readyState')
            time.sleep(.1)
        browser.assert_clean()
    with sqlite3.connect(database) as connection:
        assert connection.execute('SELECT name FROM products WHERE id=41').fetchone()[0] == 'Persisted DOM Product'
    with isolated_server(tmp_path / 'second', database) as (url, same_database), Browser(tmp_path / 'second-browser', cdn_assets) as browser:
        assert same_database == database
        browser.navigate(url + '/')
        assert_products(browser, {'p41': 'Persisted DOM Product', 'p73': 'محصول ماندگار Beta'})
        browser.assert_clean()


def test_legacy_schema_startup_backup_preserves_products(tmp_path, cdn_assets):
    with isolated_server(tmp_path / 'legacy', legacy=True) as (url, database), Browser(tmp_path / 'legacy-browser', cdn_assets) as browser:
        # Browser '/' is the first HTTP request: actual before_request must migrate.
        browser.navigate(url + '/')
        assert_products(browser)
        actual = json.loads(read_url(url + '/api/data'))
        assert actual['nodes']['p41']['quantity'] == 0
        assert actual['nodes']['a51']['required_quantity'] == 1
        backups = list(database.parent.glob('*.inventory-backup-*.sqlite'))
        assert len(backups) == 1, 'Startup must make exactly one pre-migration backup'
        with sqlite3.connect(backups[0]) as backup:
            assert 'quantity' not in {row[1] for row in backup.execute('PRAGMA table_info(products)')}
            assert backup.execute('SELECT id,name FROM products ORDER BY id').fetchall() == [(41, 'Startup Product Alpha'), (73, 'محصول ماندگار Beta')]
        with sqlite3.connect(database) as connection:
            assert 'quantity' in {row[1] for row in connection.execute('PRAGMA table_info(products)')}
            assert connection.execute('SELECT id,name FROM products ORDER BY id').fetchall() == [(41, 'Startup Product Alpha'), (73, 'محصول ماندگار Beta')]
        browser.assert_clean()
    with isolated_server(tmp_path / 'legacy-restart', database) as (url, _), Browser(tmp_path / 'legacy-restart-browser', cdn_assets) as browser:
        browser.navigate(url + '/')
        assert_products(browser)
        assert len(list(database.parent.glob('*.inventory-backup-*.sqlite'))) == 1
        browser.assert_clean()


@pytest.mark.parametrize('failure_mode', ['fail_once', 'malformed_once'])
def test_api_data_failure_is_visible_and_retry_recovers(tmp_path, cdn_assets, failure_mode):
    with isolated_server(tmp_path / 'outage') as (url, database), Browser(tmp_path / 'outage-browser', cdn_assets) as browser:
        browser.data_mode = failure_mode
        browser.navigate(url + '/')
        # Require an actual visible retry control; invoking loadData() directly would
        # hide a broken startup/error UX and is intentionally forbidden here.
        retry_expression = "Array.from(document.querySelectorAll('button')).find(e => (" + VISIBLE + ")(e) && /retry|try again|تلاش مجدد|دوباره/i.test(e.textContent))"
        browser.wait('!!(' + retry_expression + ')')
        panel_text = browser.evaluate("document.querySelector('.container').innerText")
        assert re.search(r'error|failed|unable|خطا|ناموفق', panel_text, re.I), f'No visible load failure: {panel_text}'
        assert 'هیچ محصولی تعریف نشده' not in panel_text and 'No products defined' not in panel_text
        assert not browser.evaluate("Boolean(document.querySelector('#p41_anchor'))"), 'Outage test did not intercept initial data'
        assert browser.failed_data == 1
        browser.screenshot('visible-data-error')
        browser.evaluate('(' + retry_expression + ').click(); true')
        assert_products(browser)
        assert browser.data_requests >= 2, 'Retry did not request the actual backend'
        assert not browser.evaluate('!!(' + retry_expression + ')'), 'Error retry state remained visible after recovery'
        select_inventory(browser, 'a51', 4, 3, 24, 20)
        browser.assert_clean(expected_data_failure=True)


def dom_click(browser, selector):
    browser.wait(f"({VISIBLE})(document.querySelector({json.dumps(selector)}))")
    browser.evaluate(f"document.querySelector({json.dumps(selector)}).click(); true")


def dom_fill(browser, selector, value):
    browser.wait(f"({VISIBLE})(document.querySelector({json.dumps(selector)}))")
    browser.evaluate(f"(() => {{const e=document.querySelector({json.dumps(selector)}); e.value={json.dumps(str(value))}; e.dispatchEvent(new Event('input',{{bubbles:true}})); e.dispatchEvent(new Event('change',{{bubbles:true}})); return true;}})()")


def api_until(browser, url, predicate, timeout=25):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        value = json.loads(read_url(url))
        if predicate(value):
            return value
        browser.evaluate('document.readyState')  # Service CDP interception during real HTTP polling.
        time.sleep(.1)
    raise AssertionError(f'Actual API persistence timed out: {url}; last={value}')


def directory_start(browser, url):
    browser.navigate(url + '/manufacturers')
    browser.wait("document.getElementById('directory-count').textContent.trim() === '0'")
    browser.wait("document.getElementById('directory-category').options.length > 1")
    dom_click(browser, '#directory-lang')
    browser.wait("document.documentElement.lang === 'en'")


def taxonomy_tag(browser, name):
    dom_click(browser, '#directory-taxonomy')
    browser.wait("document.querySelector('.mfg-overlay input[placeholder=\"New tag\"]')")
    dom_fill(browser, '.mfg-overlay input[placeholder="New tag"]', name)
    browser.evaluate("document.querySelector('.mfg-overlay input[placeholder=\"New tag\"]').closest('form').querySelector('button').click(); true")
    browser.wait(f"Array.from(document.querySelectorAll('.mfg-overlay input')).some(e=>e.value==={json.dumps(name)} && e.placeholder==='Name')")
    dom_click(browser, '.mfg-overlay .mfg-modal-head button')
    browser.wait(f"document.getElementById('directory-tags').textContent.includes({json.dumps(name)})")


def create_manufacturer(browser, url, name, tags=(), rich=False):
    dom_click(browser, '#directory-new')
    dom_fill(browser, '.mfg-form > label:nth-of-type(1) input', name)
    dom_fill(browser, '.mfg-form > label:nth-of-type(2) input', 'Tehran workshop 12')
    dom_fill(browser, '.mfg-form textarea', 'Audit DOM contacts')
    if rich:
        for key, values in [('phone', ['+98 21 555 0101', '+98 21 555 0102']),
                            ('email', ['sales@example.test', 'quality@example.test'])]:
            browser.evaluate(f"document.querySelector('.mfg-form input[data-key={key}]').closest('fieldset').querySelector(':scope > button').click(); true")
            for index, value in enumerate(values):
                dom_fill(browser, f'.mfg-form fieldset:nth-of-type({1 if key == "phone" else 2}) .mfg-contact-row:nth-child({index + 1}) input', value)
        dom_fill(browser, '.mfg-form input[data-key="platform"]', 'LinkedIn')
        dom_fill(browser, '.mfg-form input[data-key="handle"]', 'audit-workshop')
    for tag in tags:
        browser.evaluate(f"Array.from(document.querySelectorAll('.mfg-form .mfg-tag-choices label')).find(e=>e.textContent.endsWith({json.dumps(tag)})).querySelector('input').click(); true")
    dom_click(browser, '.mfg-form button[type="submit"]')
    browser.wait("!document.querySelector('.mfg-overlay')")
    result = api_until(browser, url + '/api/v2/manufacturers?active=all',
                       lambda result: any(m['name'] == name for m in result['data']))
    return next(m for m in result['data'] if m['name'] == name)


def export_workbook(browser, selector, expected_ids):
    from openpyxl import load_workbook
    start = len(browser.paused_exports)
    dom_click(browser, selector)
    deadline = time.monotonic() + 20
    while len(browser.paused_exports) == start:
        assert time.monotonic() < deadline, 'DOM export did not return an actual HTTP response'
        browser.evaluate('document.readyState')
        time.sleep(.08)
    response = browser.paused_exports[-1]
    assert response['responseStatusCode'] == 200
    headers = {h['name'].lower(): h['value'] for h in response['responseHeaders']}
    assert 'spreadsheetml' in headers['content-type']
    body = browser.call('Fetch.getResponseBody', {'requestId': response['requestId']})
    # Forward the unchanged real response so the app still completes its download.
    browser.call('Fetch.continueRequest', {'requestId': response['requestId']})
    payload = base64.b64decode(body['body']) if body['base64Encoded'] else body['body'].encode()
    assert payload.startswith(b'PK'), 'Export was not a real XLSX zip'
    wb = load_workbook(BytesIO(payload))
    assert wb.sheetnames == ['سازندگان', 'اطلاعات تماس', 'برچسب‌ها', 'مراحل استفاده']
    rows = list(wb.worksheets[0].values)[1:]
    assert {row[0] for row in rows} == set(expected_ids), rows
    return wb


def test_dark_manufacturer_editor_taxonomy_picker_and_settings(tmp_path, cdn_assets):
    with isolated_server(tmp_path / 'contrast-server') as (url, _), Browser(tmp_path / 'contrast-browser', cdn_assets) as browser:
        directory_start(browser, url)
        switch_theme(browser, 'dark', '#directory-theme')
        assert_style(browser, '#directory-search', '--control-text', '--control-bg')
        assert_style(browser, '#directory-search', '--placeholder-color', pseudo='::placeholder')
        for selector in ['#directory-active', '#directory-category']:
            assert_style(browser, selector, '--control-text', '--control-bg')
            assert_style(browser, selector + ' option', '--control-text', '--control-bg')
        dom_click(browser, '#directory-taxonomy')
        browser.wait("document.querySelector('.mfg-overlay input[placeholder=\"New tag\"]')")
        assert_style(browser, '.mfg-overlay input[placeholder="New tag"]', '--control-text', '--control-bg')
        assert_style(browser, '.mfg-overlay input[placeholder="New tag"]', '--placeholder-color', pseudo='::placeholder')
        assert_style(browser, '.mfg-overlay select', '--control-text', '--control-bg')
        assert_style(browser, '.mfg-overlay select option', '--control-text', '--control-bg')
        browser.screenshot('dark-manufacturer-taxonomy')
        dom_click(browser, '.mfg-overlay .mfg-modal-head button')
        create_manufacturer(browser, url, 'Contrast Workshop', rich=True)
        dom_click(browser, '#directory-new')
        browser.wait("document.querySelector('.mfg-form input[data-key=phone]')")
        for selector in ['.mfg-form > label input', '.mfg-form textarea', '.mfg-form input[data-key="phone"]']:
            assert_style(browser, selector, '--control-text', '--control-bg')
        assert_style(browser, '.mfg-form input[data-key="phone"]', '--placeholder-color', pseudo='::placeholder')
        browser.screenshot('dark-manufacturer-editor')
        dom_click(browser, '.mfg-overlay .mfg-modal-head button')
        browser.assert_clean(script='manufacturers.js')
        open_part_stages(browser, url)
        browser.wait("document.documentElement.dataset.theme === 'dark'")
        expand_stage(browser, 0)
        dom_click(browser, '.stage-item-new[data-index="0"] button[onclick="showStageManufacturerPicker(0)"]')
        browser.wait("Array.from(document.querySelectorAll('.mfg-picker-result strong')).some(e=>e.textContent==='Contrast Workshop')")
        for selector in ['.mfg-picker-filters input[type="search"]', '.mfg-picker-filters select']:
            assert_style(browser, selector, '--control-text', '--control-bg')
        assert_style(browser, '.mfg-picker-filters input[type="search"]', '--placeholder-color', pseudo='::placeholder')
        assert_style(browser, '.mfg-picker-filters select option', '--control-text', '--control-bg')
        assert_style(browser, '.mfg-picker-result strong', '--text-primary')
        browser.screenshot('dark-stage-manufacturer-picker')
        dom_click(browser, '.mfg-overlay .mfg-modal-head button')
        browser.assert_clean()
        requests = browser.data_requests
        dom_click(browser, '#settings-link')
        browser.wait("document.readyState === 'complete' && document.getElementById('theme-select')?.value === 'dark'")
        assert browser.evaluate("document.getElementById('account-role').textContent.trim()") == 'engineer'
        for selector in ['#smtp-server-input', '#theme-select', '#font-size-select']:
            assert_style(browser, selector, '--control-text', '--control-bg')
        assert_style(browser, '#smtp-server-input', '--placeholder-color', pseudo='::placeholder')
        assert_style(browser, '#theme-select option', '--control-text', '--control-bg')
        browser.screenshot('dark-settings')
        dom_fill(browser, '#theme-select', 'light')
        browser.wait("getComputedStyle(document.documentElement).colorScheme === 'light'")
        assert_style(browser, '#theme-select', '--control-text', '--control-bg')
        dom_fill(browser, '#lang-select', 'en')
        browser.wait("document.getElementById('settings-title').textContent.includes('System Settings')")
        # Shared classic script must not initialize absent jsTree on settings.
        assert browser.data_requests == requests, 'Settings unexpectedly initialized BOM /api/data'
        browser.assert_clean()


def test_manufacturer_dashboard_dom_crud_and_real_exports(tmp_path, cdn_assets):
    with isolated_server(tmp_path / 'directory') as (url, database), Browser(tmp_path / 'directory-browser', cdn_assets) as browser:
        directory_start(browser, url)
        taxonomy_tag(browser, 'Audit precision')
        taxonomy_tag(browser, 'Audit certified')
        alpha = create_manufacturer(browser, url, 'Audit Alpha', ['Audit precision', 'Audit certified'], rich=True)
        beta = create_manufacturer(browser, url, 'Audit Beta', ['Audit precision'])
        gamma = create_manufacturer(browser, url, 'Audit Gamma', ['Audit certified'])
        browser.wait("document.getElementById('directory-count').textContent === '3'")
        assert [p['phone'] for p in alpha['phones']] == ['+98 21 555 0101', '+98 21 555 0102']
        assert [e['email'] for e in alpha['emails']] == ['sales@example.test', 'quality@example.test']
        assert [(s['platform'], s['handle']) for s in alpha['socials']] == [('LinkedIn', 'audit-workshop')]
        browser.evaluate("Array.from(document.querySelectorAll('.mfg-directory-card')).find(e=>e.querySelector('h2').textContent==='Audit Alpha').querySelector('button').click(); true")
        dom_fill(browser, '.mfg-form > label:nth-of-type(1) input', 'Audit Alpha edited')
        dom_fill(browser, '.mfg-form input[data-key="phone"]', '+98 21 555 0199')
        dom_click(browser, '.mfg-form button[type="submit"]')
        browser.wait("!document.querySelector('.mfg-overlay')")
        edited = api_until(browser, url + f'/api/v2/manufacturers/{alpha["id"]}', lambda r: r['data']['name'] == 'Audit Alpha edited')['data']
        assert edited['id'] == alpha['id']
        assert [p['phone'] for p in edited['phones']] == ['+98 21 555 0199', '+98 21 555 0102']
        assert len(edited['emails']) == 2 and len(edited['tags']) == 2
        browser.evaluate("Array.from(document.querySelectorAll('#directory-tags label')).find(e=>e.textContent.endsWith('Audit precision')).querySelector('input').click(); true")
        browser.wait("document.getElementById('directory-count').textContent === '2'")
        browser.evaluate("Array.from(document.querySelectorAll('#directory-tags label')).find(e=>e.textContent.endsWith('Audit certified')).querySelector('input').click(); true")
        browser.wait("document.getElementById('directory-count').textContent === '1'")
        assert browser.evaluate("Array.from(document.querySelectorAll('.mfg-directory-card h2'),e=>e.textContent)") == ['Audit Alpha edited'], 'Tag filters must use AND, not OR'
        filtered = export_workbook(browser, '#directory-export-filtered', [alpha['id']])
        contacts = list(filtered['اطلاعات تماس'].values)[1:]
        assert {r[3] for r in contacts} == {'+98 21 555 0199', '+98 21 555 0102', 'sales@example.test', 'quality@example.test', 'audit-workshop'}
        assert {r[3] for r in list(filtered['برچسب‌ها'].values)[1:]} == {'Audit precision', 'Audit certified'}
        export_workbook(browser, '.mfg-directory-card .mfg-actions button:nth-child(2)', [alpha['id']])
        export_workbook(browser, '#directory-export-all', [alpha['id'], beta['id'], gamma['id']])
        browser.screenshot('manufacturer-dashboard-workflow')
        browser.assert_clean(script='manufacturers.js')
        with sqlite3.connect(database) as connection:
            assert connection.execute('SELECT count(*) FROM manufacturers').fetchone()[0] == 3
            assert connection.execute('SELECT count(*) FROM part_manufacturers').fetchone()[0] == 0


def open_part_stages(browser, url):
    browser.navigate(url + '/')
    assert_products(browser)
    # Expand by visible jsTree toggles instead of bypassing the widget API.
    for node in ['p41', 'a51', 'a52']:
        browser.wait(f"document.getElementById('{node}')")
        if browser.evaluate(f"document.getElementById('{node}').classList.contains('jstree-closed')"):
            dom_click(browser, f'#{node} > .jstree-ocl')
    dom_click(browser, '#r61_anchor')
    dom_click(browser, '#tab-btn-mfg')
    browser.wait("document.querySelectorAll('#stages-list .stage-item-new').length >= 2")


def expand_stage(browser, index):
    selector = f'.stage-item-new[data-index="{index}"] button[aria-expanded]'
    if browser.evaluate(f"document.querySelector({json.dumps(selector)}).getAttribute('aria-expanded') === 'false'"):
        dom_click(browser, selector)
    browser.wait(f"document.getElementById('stage-manufacturers-{index}')")


def picker_add(browser, index, name, close=True):
    expand_stage(browser, index)
    dom_click(browser, f'.stage-item-new[data-index="{index}"] button[onclick="showStageManufacturerPicker({index})"]')
    browser.wait(f"Array.from(document.querySelectorAll('.mfg-picker-result strong')).some(e=>e.textContent==={json.dumps(name)})")
    browser.evaluate(f"Array.from(document.querySelectorAll('.mfg-picker-result')).find(e=>e.querySelector('strong').textContent==={json.dumps(name)}).querySelector('button').click(); true")
    browser.wait(f"document.getElementById('stage-manufacturers-{index}').textContent.includes({json.dumps(name)})")
    browser.wait(f"!Array.from(document.querySelectorAll('.mfg-picker-result strong')).some(e=>e.textContent==={json.dumps(name)})")
    if close:
        dom_click(browser, '.mfg-overlay .mfg-modal-head button')


def stage_rows(browser, index):
    return browser.evaluate(f"Array.from(document.querySelectorAll('#stage-manufacturers-{index} .mfg-stage-row > strong'),e=>e.textContent)")


def test_stage_manufacturer_distinct_lists_order_and_persistence(tmp_path, cdn_assets):
    # Keep fixture DB outside the repository, even with an in-workspace --basetemp.
    with tempfile.TemporaryDirectory(prefix='bom-stage-browser-') as server_dir, \
         isolated_server(Path(server_dir)) as (url, database), \
         Browser(tmp_path / 'stage-browser', cdn_assets) as browser:
        directory_start(browser, url)
        alpha = create_manufacturer(browser, url, 'Stage Alpha', rich=True)
        beta = create_manufacturer(browser, url, 'Stage Beta')
        gamma = create_manufacturer(browser, url, 'Stage Gamma')
        open_part_stages(browser, url)
        picker_add(browser, 0, 'Stage Alpha')
        picker_add(browser, 0, 'Stage Beta')
        picker_add(browser, 1, 'Stage Gamma')
        expand_stage(browser, 0)
        expand_stage(browser, 1)
        assert stage_rows(browser, 0) == ['1. Stage Alpha', '2. Stage Beta']
        assert stage_rows(browser, 1) == ['1. Stage Gamma']
        duplicate_ids = browser.evaluate("(() => {const ids=Array.from(document.querySelectorAll('[id]'),e=>e.id); return ids.filter((id,i)=>ids.indexOf(id)!==i);})()")
        assert not duplicate_ids, f'Duplicate DOM IDs corrupt simultaneous stage panels: {duplicate_ids}'
        dom_click(browser, '#stage-manufacturers-0 .mfg-stage-row:nth-child(2) button[aria-label="Move up"]')
        assert stage_rows(browser, 0) == ['1. Stage Beta', '2. Stage Alpha']
        result = api_until(browser, url + '/api/data', lambda d: d['nodes']['r61']['stages'][0]['manufacturer_ids'] == [beta['id'], alpha['id']] and d['nodes']['r61']['stages'][1]['manufacturer_ids'] == [gamma['id']])
        assert [s['id'] for s in result['nodes']['r61']['stages']] == [81, 82]
        with closing(sqlite3.connect(database)) as connection:
            assert connection.execute('SELECT id,stage_id,description FROM stage_details WHERE id=91').fetchone() == (91, 81, 'Preserve fixture machining detail')
            assert connection.execute('SELECT count(*) FROM part_manufacturers').fetchone()[0] == 0
        open_part_stages(browser, url)
        expand_stage(browser, 0)
        expand_stage(browser, 1)
        browser.wait("document.getElementById('stage-details-0').textContent.includes('Preserve fixture machining detail')")
        assert stage_rows(browser, 0) == ['1. Stage Beta', '2. Stage Alpha']
        assert stage_rows(browser, 1) == ['1. Stage Gamma']
        picker_add(browser, 1, 'Stage Alpha', close=False)
        picker_chips = browser.evaluate("""(() => {
            return Array.from(document.querySelectorAll('.mfg-overlay .mfg-picker-filters .mfg-tag-choices label'), label => {
                const checkbox = label.querySelector('input[type="checkbox"]');
                const bounds = label.getBoundingClientRect();
                return {text: label.textContent.trim(), height: bounds.height,
                        checkboxWidth: checkbox ? checkbox.getBoundingClientRect().width : null};
            });
        })()""")
        assert picker_chips, 'Open picker must expose tag chips for the layout regression check'
        for chip in picker_chips:
            assert chip['text'] and 0 < chip['height'] <= 40, f'Tag chip wrapped into a narrow vertical column: {chip}'
            assert chip['checkboxWidth'] is not None and 0 < chip['checkboxWidth'] <= 25, f'Picker checkbox inherited full-width input styling: {chip}'
        browser.screenshot('stage-manufacturer-picker')
        dom_click(browser, '.mfg-overlay .mfg-modal-head button')
        api_until(browser, url + '/api/data', lambda d: d['nodes']['r61']['stages'][1]['manufacturer_ids'] == [gamma['id'], alpha['id']])
        browser.assert_clean()


def test_new_stage_second_edit_during_delayed_first_save_keeps_id(tmp_path, cdn_assets):
    with isolated_server(tmp_path / 'delayed') as (url, database), Browser(tmp_path / 'delayed-browser', cdn_assets) as browser:
        directory_start(browser, url)
        alpha = create_manufacturer(browser, url, 'Delayed Alpha')
        beta = create_manufacturer(browser, url, 'Delayed Beta')
        gamma = create_manufacturer(browser, url, 'Delayed Gamma')
        open_part_stages(browser, url)
        browser.hold_node_save = True
        dom_fill(browser, '#new-stage-name', 'New delayed stage')
        dom_click(browser, '.stage-add button')
        browser.wait("document.querySelectorAll('#stages-list .stage-item-new').length === 3")
        deadline = time.monotonic() + 10
        while not browser.paused_node_saves:
            assert time.monotonic() < deadline, 'First new-stage autosave was never sent'
            browser.evaluate('document.readyState')
            time.sleep(.08)
        first_request = browser.paused_node_saves[0]
        first_payload = json.loads(first_request['request']['postData'])
        assert first_payload['stages'][2]['name'] == 'New delayed stage'
        assert not first_payload['stages'][2].get('id'), 'Test must exercise a genuinely unpersisted stage'
        picker_add(browser, 2, 'Delayed Alpha')
        dom_fill(browser, '.stage-item-new[data-index="2"] input[type="number"]', 1234)
        # Leave picker open over identity reconciliation; its context must still resolve.
        picker_add(browser, 2, 'Delayed Beta', close=False)
        # Ensure the second snapshot is queued while the first HTTP request is held.
        end = time.monotonic() + .9
        while time.monotonic() < end:
            browser.evaluate('document.readyState')
            time.sleep(.08)
        browser.hold_node_save = False
        for paused in browser.paused_node_saves:
            browser.call('Fetch.continueRequest', {'requestId': paused['requestId']})
        browser.paused_node_saves.clear()
        first = api_until(browser, url + '/api/data', lambda d: len(d['nodes']['r61']['stages']) == 3)
        created_id = first['nodes']['r61']['stages'][2]['id']
        final = api_until(browser, url + '/api/data', lambda d: len(d['nodes']['r61']['stages']) == 3 and d['nodes']['r61']['stages'][2]['manufacturer_ids'] == [alpha['id'], beta['id']] and d['nodes']['r61']['stages'][2]['estimated_material_cost'] == 1234)
        assert final['nodes']['r61']['stages'][2]['id'] == created_id, 'Queued temporary stage edit recreated the persisted stage ID'
        saves = [e['params']['request'] for e in browser.events if e['method'] == 'Network.requestWillBeSent'
                 and e['params']['request']['method'] == 'PUT' and '/api/node/r61' in e['params']['request']['url']]
        assert len(saves) >= 2, 'Second edit never produced a queued real HTTP save'
        assert json.loads(saves[-1]['postData'])['stages'][2]['id'] == created_id, 'Queued save did not reconcile its temporary stage to the first persisted ID'
        browser.wait("document.querySelector('.mfg-overlay')")
        browser.wait("Array.from(document.querySelectorAll('.mfg-picker-result strong')).some(e=>e.textContent==='Delayed Gamma')")
        browser.evaluate("Array.from(document.querySelectorAll('.mfg-picker-result')).find(e=>e.querySelector('strong').textContent==='Delayed Gamma').querySelector('button').click(); true")
        after_picker = api_until(browser, url + '/api/data', lambda d: d['nodes']['r61']['stages'][2]['manufacturer_ids'] == [alpha['id'], beta['id'], gamma['id']])
        assert after_picker['nodes']['r61']['stages'][2]['id'] == created_id, 'Open picker lost its context after stage identity reconciliation'
        dom_click(browser, '.mfg-overlay .mfg-modal-head button')
        with sqlite3.connect(database) as connection:
            assert connection.execute('SELECT count(*) FROM stages WHERE part_id=61').fetchone()[0] == 3
            assert connection.execute('SELECT stage_id FROM stage_details WHERE id=91').fetchone()[0] == 81
        open_part_stages(browser, url)
        expand_stage(browser, 2)
        assert stage_rows(browser, 2) == ['1. Delayed Alpha', '2. Delayed Beta', '3. Delayed Gamma']
        assert browser.evaluate("document.querySelector('.stage-item-new[data-index=\"2\"] input[type=number]').value") == '1234'
        browser.assert_clean()
