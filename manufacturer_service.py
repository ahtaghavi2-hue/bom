"""Validation, filtering and exports shared by manufacturer and stage endpoints."""
from io import BytesIO
from sqlalchemy import or_
from werkzeug.exceptions import BadRequest
from extensions import db
from models import (Manufacturer, ManufacturerEmail, ManufacturerPhone, ManufacturerSocial,
                    Category, Tag, StageManufacturer)


def json_object(request):
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        raise BadRequest('JSON object required')
    return data


def validate_ids(value, model, field):
    if not isinstance(value, list) or any(type(i) is not int or i <= 0 for i in value):
        raise ValueError(field + ' must be an array of positive integer IDs')
    if len(set(value)) != len(value):
        raise ValueError(field + ' must be unique')
    if value and model.query.filter(model.id.in_(value)).count() != len(value):
        raise ValueError(field + ' contains unknown IDs')
    return value


def validate_manufacturer(data, creating=False):
    result = {}
    for key, limit in (('name', 200), ('phone', 50), ('address', None), ('notes', None)):
        if key not in data:
            continue
        value = data[key]
        if not isinstance(value, str) or (limit and len(value) > limit) or (key == 'name' and not value.strip()):
            raise ValueError('Invalid ' + key)
        result[key] = value.strip() if key == 'name' else value
    if creating and 'name' not in result:
        result['name'] = 'سازنده جدید'
    if 'is_active' in data:
        if type(data['is_active']) is not bool:
            raise ValueError('is_active must be boolean')
        result['is_active'] = data['is_active']
    if 'tag_ids' in data:
        result['tag_ids'] = validate_ids(data['tag_ids'], Tag, 'tag_ids')
    for key, fields in (('emails', {'email': 200}), ('phones', {'phone': 50}),
                        ('socials', {'platform': 100, 'handle': 200})):
        if key not in data:
            continue
        if not isinstance(data[key], list):
            raise ValueError(key + ' must be an array of objects')
        rows, seen = [], set()
        for item in data[key]:
            if not isinstance(item, dict):
                raise ValueError('Invalid ' + key + ' item')
            row = {}
            for field, limit in fields.items():
                value = item.get(field)
                if not isinstance(value, str) or not value.strip() or len(value.strip()) > limit:
                    raise ValueError('Invalid ' + key + '.' + field)
                row[field] = value.strip()
            identity = tuple(row.values())
            if identity not in seen:
                rows.append(row)
                seen.add(identity)
        result[key] = rows
    return result


def apply_manufacturer(manufacturer, values):
    for key in ('name', 'phone', 'address', 'notes', 'is_active'):
        if key in values:
            setattr(manufacturer, key, values[key])
    if 'tag_ids' in values:
        manufacturer.tags = Tag.query.filter(Tag.id.in_(values['tag_ids'])).order_by(Tag.id).all()
    for key, model in (('emails', ManufacturerEmail), ('phones', ManufacturerPhone), ('socials', ManufacturerSocial)):
        if key in values:
            relationship = getattr(manufacturer, key)
            for old in relationship.all():
                db.session.delete(old)
            for row in values[key]:
                relationship.append(model(**row))


def replace_stage_manufacturers(stage, ids):
    # Reuse associations so reorders cannot collide with the composite primary key.
    existing = {link.manufacturer_id: link for link in stage.manufacturer_links}
    stage.manufacturer_links[:] = [existing.get(mid) or StageManufacturer(manufacturer_id=mid)
                                  for mid in ids]
    for priority, link in enumerate(stage.manufacturer_links, 1):
        link.priority = priority


def stage_manufacturers(stage):
    links = sorted(stage.manufacturer_links, key=lambda link: link.priority)
    return {'manufacturer_ids': [link.manufacturer_id for link in links],
            'manufacturers': [dict(link.manufacturer.to_dict(), priority=link.priority) for link in links]}


def manufacturer_usage(manufacturer):
    result = []
    for link in sorted(manufacturer.stage_links, key=lambda item: (item.stage_id, item.priority)):
        stage = link.stage
        part = stage.part
        product = part.assembly.product
        result.append({'product_id': product.id, 'product_name': product.name,
                       'part_id': part.id, 'part_name': part.name,
                       'stage_id': stage.id, 'stage_name': stage.name, 'priority': link.priority})
    return result


def filter_manufacturers(args):
    query = Manufacturer.query
    active = args.get('active', 'all')
    if active not in ('all', 'true', 'false'):
        raise ValueError('active must be all, true or false')
    if active != 'all':
        query = query.filter(Manufacturer.is_active == (active == 'true'))
    def parameter_id(raw):
        if not raw.isdecimal() or int(raw) <= 0:
            raise ValueError('Invalid filter ID')
        return int(raw)
    for tag_id in args.getlist('tag_id'):
        query = query.filter(Manufacturer.tags.any(Tag.id == parameter_id(tag_id)))
    if args.get('category_id') is not None:
        query = query.filter(Manufacturer.tags.any(Tag.category_id == parameter_id(args['category_id'])))
    q = args.get('q', '').strip()
    if q:
        pattern = '%' + q.replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_') + '%'
        def contains(column):
            return column.ilike(pattern, escape='\\')
        query = query.filter(or_(contains(Manufacturer.name), contains(Manufacturer.phone),
            contains(Manufacturer.address), contains(Manufacturer.notes),
            Manufacturer.emails.any(contains(ManufacturerEmail.email)),
            Manufacturer.phones.any(contains(ManufacturerPhone.phone)),
            Manufacturer.socials.any(or_(contains(ManufacturerSocial.platform), contains(ManufacturerSocial.handle))),
            Manufacturer.tags.any(contains(Tag.name))))
    return query.order_by(Manufacturer.name, Manufacturer.id)


def manufacturer_workbook(manufacturers):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    def sheet(title, headers, rows):
        ws = wb.create_sheet(title)
        for row_index, row in enumerate([headers, *rows], 1):
            for column, value in enumerate(row, 1):
                cell = ws.cell(row_index, column, value)
                if isinstance(value, str):
                    cell.data_type = 's'
                    cell.number_format = '@'
                cell.font = Font(name='Arial', bold=row_index == 1)
                cell.alignment = Alignment(horizontal='right', wrap_text=True)
                if row_index == 1:
                    cell.fill = PatternFill('solid', fgColor='DCE6F1')
        ws.sheet_view.rightToLeft = True
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for column in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(column)].width = 24
    summaries, contacts, tags, usages = [], [], [], []
    for manufacturer in manufacturers:
        data = manufacturer.to_dict()
        prefix = [manufacturer.id, manufacturer.name]
        summaries.append(prefix + ['فعال' if manufacturer.is_active else 'غیرفعال', manufacturer.phone or '',
                                   manufacturer.address or '', manufacturer.notes or ''])
        for key, label, field in (('phones', 'تلفن', 'phone'), ('emails', 'ایمیل', 'email')):
            for contact in data[key]:
                contacts.append(prefix + [label, contact[field], ''])
        for social in data['socials']:
            contacts.append(prefix + ['شبکه اجتماعی', social['handle'], social['platform']])
        for tag in data['tags']:
            tags.append(prefix + [tag['category_name'], tag['name']])
        for usage in manufacturer_usage(manufacturer):
            usages.append(prefix + [usage['product_name'], usage['part_name'], usage['stage_name'], usage['priority']])
    sheet('سازندگان', ['شناسه', 'نام', 'وضعیت', 'تلفن قدیمی', 'آدرس', 'یادداشت'], summaries)
    sheet('اطلاعات تماس', ['شناسه سازنده', 'نام سازنده', 'نوع', 'مقدار', 'پلتفرم'], contacts)
    sheet('برچسب‌ها', ['شناسه سازنده', 'نام سازنده', 'دسته‌بندی', 'برچسب'], tags)
    sheet('مراحل استفاده', ['شناسه سازنده', 'نام سازنده', 'محصول', 'قطعه', 'مرحله', 'اولویت'], usages)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
