"""Regression tests: exclusively in-memory or pytest temporary SQLite files."""
from io import BytesIO
from pathlib import Path
import sqlite3
from unittest.mock import patch

import pytest
from sqlalchemy import create_engine, inspect, text

from config import Config

# Override configuration BEFORE app import/extension initialization. No real DB
# is opened, migrated, or dropped by this module.
with patch.object(Config, 'SQLALCHEMY_DATABASE_URI', 'sqlite:///:memory:'):
    import app as backend

from extensions import db
from models import Product, Assembly, Part, Stage, StageDetail, ProductionSchedule, WorkOrder, User
from bom_inventory_migration import ensure_inventory_schema


@pytest.fixture
def client():
    app = backend.app
    config = dict(app.config)
    keys = ('bom_inventory_schema_ready', 'manufacturer_schema_ready')
    flags = {key: app.extensions[key] for key in keys if key in app.extensions}
    try:
        app.config.update(TESTING=True, LOGIN_DISABLED=False)
        with app.app_context():
            assert db.engine.url.database == ':memory:'
            try:
                db.create_all()
                with db.engine.begin() as conn:
                    conn.execute(text('DROP TABLE IF EXISTS manufacturer_schema_state'))
                for key in keys:
                    app.extensions.pop(key, None)
                # Classic mutation routes need a real role now, so sign in an admin.
                admin = User(username='admin', email='admin@example.com', role='admin')
                admin.set_password('test-password')
                db.session.add(admin)
                db.session.commit()
                client = app.test_client()
                with client.session_transaction() as session:
                    session['_user_id'] = str(admin.id)
                    session['_fresh'] = True
                yield client
            finally:
                db.session.remove()
                db.drop_all()
                with db.engine.begin() as conn:
                    conn.execute(text('DROP TABLE IF EXISTS manufacturer_schema_state'))
    finally:
        app.config.clear()
        app.config.update(config)
        for key in keys:
            app.extensions.pop(key, None)
        app.extensions.update(flags)


def create(client, typ, **fields):
    response = client.post('/api/node', json={'type': typ, 'name': typ, **fields})
    assert response.status_code == 200, response.get_data(as_text=True)
    return response.json['node']


def tree(client):
    product = create(client, 'product', name='Product', order_count=10, quantity=2)
    assembly = create(client, 'assembly', name='Assembly', parent=product['id'], required_quantity=3, quantity=5)
    nested = create(client, 'assembly', name='Nested', parent=assembly['id'], required_quantity=2, quantity=7)
    part = create(client, 'part', name='Part', parent=nested['id'], required_quantity=4, quantity=6,
                  partCode='SAME', supplier_email='supplier@example.com')
    sibling = create(client, 'part', name='Sibling', parent=nested['id'], required_quantity=4, quantity=200,
                     partCode='SAME')
    return product, assembly, nested, part, sibling


def test_recursive_stock_and_independent_siblings(client):
    original = tree(client)
    nodes = client.get('/api/data').json['nodes']
    expected = [(10, 8, 2), (24, 19, 5), (38, 31, 7), (124, 118, 6), (124, 0, 124)]
    for item, values in zip(original, expected):
        node = nodes[item['id']]
        assert (node['required_total'], node['shortage'], node['stock_used']) == values
        assert node['quantity'] == item['quantity']
    assert client.get('/api/data').json['nodes'] == nodes
    shortages = client.get('/api/notifications').json['shortage']
    assert shortages == [{'id': original[3]['id'], 'name': 'Part', 'required': 124,
                          'available': 6, 'shortage': 118}]
    assert client.get('/api/v2/dashboard').json['data']['low_stock_parts'] == 1


def test_zero_demand_and_zero_coefficients(client):
    product, assembly, nested, part, sibling = tree(client)
    assert client.patch('/api/node/' + product['id'], json={'order_count': 0}).status_code == 200
    for node in client.get('/api/data').json['nodes'].values():
        assert node['required_total'] == node['shortage'] == node['stock_used'] == 0
    client.patch('/api/node/' + product['id'], json={'order_count': 10})
    client.patch('/api/node/' + assembly['id'], json={'required_quantity': 0})
    nodes = client.get('/api/data').json['nodes']
    assert nodes[assembly['id']]['required_quantity'] == 0
    assert all(nodes[n['id']]['required_total'] == 0 for n in (assembly, nested, part, sibling))


def test_roundtrip_mapped_fields_and_stage_details(client):
    product = create(client, 'product', order_count='4', quantity='1')
    assembly = create(client, 'assembly', parent=product['id'], quantity='2', required_quantity='3')
    part = create(client, 'part', parent=assembly['id'], quantity='5', required_quantity='6',
                  partCode='CODE', partType='buy', supplier_email='first@example.com',
                  stages=[{'name': 'Cut', 'estimated_material_cost': 25}])
    assert part['id'].startswith('r')
    model = db.session.get(Part, int(part['id'][1:]))
    stage = model.stages.first()
    stage_id = stage.id
    detail = StageDetail(stage_id=stage_id, description='Keep this')
    db.session.add(detail)
    schedule = ProductionSchedule(product_id=int(product['id'][1:]), quantity=99)
    db.session.add(schedule)
    db.session.flush()
    order = WorkOrder(schedule_id=schedule.id, part_id=model.id, quantity=88)
    db.session.add(order)
    db.session.commit()
    response = client.put('/api/node/' + part['id'], json=part)
    assert response.status_code == 200
    db.session.expire_all()
    assert model.stages.first().id == stage_id
    assert db.session.get(StageDetail, detail.id).description == 'Keep this'
    assert db.session.get(ProductionSchedule, schedule.id).quantity == 99
    assert db.session.get(WorkOrder, order.id).quantity == 88
    response = client.patch('/api/node/' + part['id'], json={
        'partCode': 'NEW', 'partType': 'make', 'supplier_email': 'new@example.com'})
    assert response.json['node']['partCode'] == 'NEW'
    db.session.expire_all()
    assert (model.part_code, model.part_type, model.supplier_email) == ('NEW', 'make', 'new@example.com')
    assert response.json['node']['required_total'] == 42


@pytest.fixture
def staged_part(client):
    product = create(client, 'product')
    assembly = create(client, 'assembly', parent=product['id'])
    part = create(client, 'part', parent=assembly['id'], stages=[
        {'name': name, 'estimated_material_cost': cost}
        for name, cost in [('Cut', 11), ('Drill', 22), ('Paint', 33)]])
    details = {}
    for stage in part['stages']:
        detail = StageDetail(stage_id=stage['id'], description=stage['name'] + ' detail')
        db.session.add(detail)
        db.session.flush()
        details[stage['id']] = detail.id
    db.session.commit()
    return part, details


def assert_stage_details(stages, details):
    for stage in stages:
        detail = db.session.get(StageDetail, details[stage['id']])
        assert detail.stage_id == stage['id']
        assert detail.description == stage['name'] + ' detail'


def test_stage_reorder_preserves_identity_costs_and_details(client, staged_part):
    part, details = staged_part
    reordered = [part['stages'][2], part['stages'][0], part['stages'][1]]
    response = client.patch('/api/node/' + part['id'], json={
        'stages': [{'id': stage['id'], 'name': stage['name']} for stage in reordered]})
    assert response.status_code == 200
    assert response.json['node']['stages'] == reordered
    db.session.expire_all()
    assert_stage_details(reordered, details)
    assert [db.session.get(Stage, stage['id']).sort_order for stage in reordered] == [0, 1, 2]


def test_stage_middle_deletion_preserves_survivor_details(client, staged_part):
    part, details = staged_part
    survivors = [part['stages'][0], part['stages'][2]]
    removed_id = part['stages'][1]['id']
    response = client.put('/api/node/' + part['id'], json={**part, 'stages': survivors})
    assert response.status_code == 200
    assert response.json['node']['stages'] == survivors
    db.session.expire_all()
    assert_stage_details(survivors, details)
    assert db.session.get(Stage, removed_id) is None
    assert db.session.get(StageDetail, details[removed_id]) is None
    assert db.session.get(Stage, survivors[1]['id']).sort_order == 1


def test_mixed_stage_ids_create_new_entries_without_positional_reuse(client, staged_part):
    part, details = staged_part
    retained = part['stages'][2]
    response = client.patch('/api/node/' + part['id'], json={'stages': [
        {'name': 'New'}, {'id': retained['id']}]})
    assert response.status_code == 200
    stages = response.json['node']['stages']
    assert stages[0]['name'] == 'New'
    assert stages[0]['id'] not in details
    assert stages[1] == retained
    db.session.expire_all()
    assert_stage_details([retained], details)
    for old_stage in part['stages'][:2]:
        assert db.session.get(Stage, old_stage['id']) is None
        assert db.session.get(StageDetail, details[old_stage['id']]) is None


def test_legacy_stage_lists_preserve_positions_and_append(client, staged_part):
    part, details = staged_part
    payload = [{'name': stage['name']} for stage in part['stages']] + [{'name': 'New'}]
    response = client.patch('/api/node/' + part['id'], json={'stages': payload})
    assert response.status_code == 200
    stages = response.json['node']['stages']
    assert stages[:3] == part['stages']
    assert stages[3]['id'] not in details
    db.session.expire_all()
    assert_stage_details(stages[:3], details)
    response = client.patch('/api/node/' + part['id'], json={'stages': []})
    assert response.status_code == 200
    assert response.json['node']['stages'] == []
    assert Stage.query.count() == StageDetail.query.count() == 0


@pytest.mark.parametrize('bad_id', ['duplicate', 'foreign', 'unknown', None, True, 1.0, '1', [], {}])
def test_invalid_stage_ids_rejected_before_mutation(client, staged_part, bad_id):
    part, details = staged_part
    other = create(client, 'part', parent=part['parent'], stages=[{'name': 'Other'}])
    if bad_id == 'duplicate':
        invalid_id = part['stages'][0]['id']
    elif bad_id == 'foreign':
        invalid_id = other['stages'][0]['id']
    elif bad_id == 'unknown':
        invalid_id = 999999
    else:
        invalid_id = bad_id
    payload = {'name': 'Wrong', 'quantity': 99, 'stages': [
        {'id': part['stages'][0]['id'], 'name': 'Wrong', 'estimated_material_cost': 999},
        {'id': invalid_id, 'name': 'Invalid'}]}
    model = db.session.get(Part, int(part['id'][1:]))
    # Check the helper itself, not just endpoint rollback: validation must
    # happen before any mapped fields, stage fields, inserts, or deletes change.
    with pytest.raises(ValueError):
        backend._apply_node_fields(model, 'part', payload)
    assert model.name == part['name'] and model.quantity == part['quantity']
    assert not db.session.dirty and not db.session.new and not db.session.deleted
    response = client.patch('/api/node/' + part['id'], json=payload)
    assert response.status_code == 400
    assert client.get('/api/data').json['nodes'][part['id']] == part
    assert client.get('/api/data').json['nodes'][other['id']] == other
    assert_stage_details(part['stages'], details)
    response = client.post('/api/node', json={
        'type': 'part', 'parent': part['parent'], 'stages': [{'id': other['stages'][0]['id']}]})
    assert response.status_code == 400
    assert Part.query.count() == 2


@pytest.mark.parametrize('bad', [-1, 1.5, True, None, '1.2', '-2', ''])
def test_invalid_inventory_is_atomic(client, bad):
    product = create(client, 'product', quantity=5)
    response = client.patch('/api/node/' + product['id'], json={'name': 'Wrong', 'quantity': bad})
    assert response.status_code == 400
    node = client.get('/api/data').json['nodes'][product['id']]
    assert node['name'] == 'product' and node['quantity'] == 5
    assert client.post('/api/node', json={'type': 'product', 'quantity': bad}).status_code == 400


def test_invalid_parents_and_moves_preserve_subtree(client):
    product, assembly, nested, part, sibling = tree(client)
    other = create(client, 'product', order_count=2)
    for target in (assembly['id'], nested['id'], part['id'], '#', 'a999', 'invalid'):
        assert client.post('/api/node/' + assembly['id'] + '/move', json={'new_parent': target}).status_code == 400
    assert client.post('/api/node/' + part['id'] + '/move', json={'new_parent': product['id']}).status_code == 400
    assert client.post('/api/node', json={'type': 'part', 'parent': product['id']}).status_code == 400
    assert client.post('/api/node', json={'type': 'assembly', 'parent': part['id']}).status_code == 400
    assert client.post('/api/node/' + assembly['id'] + '/move', json={'new_parent': other['id']}).status_code == 200
    db.session.expire_all()
    for node in (assembly, nested):
        assert db.session.get(Assembly, int(node['id'][1:])).product_id == int(other['id'][1:])
    assert client.get('/api/data').json['nodes'][part['id']]['required_total'] == 0


def test_serializers_bind_fields_per_model(client):
    product, assembly, nested, part, sibling = tree(client)
    p = db.session.get(Product, int(product['id'][1:]))
    a = db.session.get(Assembly, int(assembly['id'][1:]))
    r = db.session.get(Part, int(part['id'][1:]))
    assert p.to_dict()['quantity'] == 2
    assert a.to_dict()['required_quantity'] == 3
    assert r.to_dict()['part_code'] == 'SAME'
    assert 'description' in p.to_dict(full=True)
    assert 'supplier_email' in r.to_dict(full=True)
    assert 'schedule_id' not in p.to_dict()
    assert client.get('/api/v2/products').status_code == 200
    assert client.get('/api/v2/products/' + product['id'][1:] + '/assemblies').json['data'][0]['quantity'] == 5


@pytest.mark.parametrize('existing_assembly_columns', [False, True])
def test_schema_backup_idempotence_and_preservation(tmp_path, existing_assembly_columns):
    path = tmp_path / 'legacy.sqlite'
    engine = create_engine('sqlite:///' + path.as_posix())
    with engine.begin() as connection:
        connection.execute(text('CREATE TABLE products (id INTEGER PRIMARY KEY, name TEXT)'))
        extra = ', quantity INTEGER, required_quantity INTEGER' if existing_assembly_columns else ''
        connection.execute(text('CREATE TABLE assemblies (id INTEGER PRIMARY KEY, name TEXT' + extra + ')'))
        connection.execute(text('CREATE TABLE parts (id INTEGER PRIMARY KEY, quantity INTEGER, required_quantity INTEGER)'))
        connection.execute(text("INSERT INTO products VALUES (1, 'unchanged')"))
        if existing_assembly_columns:
            connection.execute(text("INSERT INTO assemblies VALUES (1, 'null', NULL, NULL), (2, 'values', 9, 0)"))
        else:
            connection.execute(text("INSERT INTO assemblies VALUES (1, 'unchanged')"))
        connection.execute(text('INSERT INTO parts VALUES (1, NULL, 7), (2, 8, NULL), (3, 5, 0)'))
        connection.execute(text('CREATE TABLE production_schedules (id INTEGER, quantity INTEGER)'))
        connection.execute(text('INSERT INTO production_schedules VALUES (1, 999)'))
    backup = ensure_inventory_schema(engine)
    assert backup and Path(backup).exists()
    with sqlite3.connect(backup) as saved:
        assert 'quantity' not in {c[1] for c in saved.execute('PRAGMA table_info(products)')}
        assert saved.execute('SELECT quantity FROM parts WHERE id=1').fetchone()[0] is None
    with engine.connect() as connection:
        assert connection.execute(text('SELECT name, quantity FROM products')).all() == [('unchanged', 0)]
        assert connection.execute(text('SELECT quantity, required_quantity FROM parts ORDER BY id')).all() == [(0, 7), (8, 1), (5, 0)]
        if existing_assembly_columns:
            assert connection.execute(text('SELECT quantity, required_quantity FROM assemblies ORDER BY id')).all() == [(0, 1), (9, 0)]
        assert connection.execute(text('SELECT quantity FROM production_schedules')).scalar() == 999
    assert ensure_inventory_schema(engine) is None
    assert len(list(tmp_path.glob('*.inventory-backup-*'))) == 1
    engine.dispose()


def test_backup_failure_prevents_schema_changes(tmp_path):
    engine = create_engine('sqlite:///' + (tmp_path / 'legacy.sqlite').as_posix())
    with engine.begin() as connection:
        connection.execute(text('CREATE TABLE products (id INTEGER PRIMARY KEY)'))
    with patch('bom_inventory_migration.sqlite3.connect', side_effect=OSError('backup failed')):
        with pytest.raises(OSError):
            ensure_inventory_schema(engine)
    assert {c['name'] for c in inspect(engine).get_columns('products')} == {'id'}
    engine.dispose()


def test_excel_nested_sheets_and_zero_demand(client):
    from openpyxl import load_workbook
    product, assembly, nested, part, sibling = tree(client)
    response = client.get('/api/export/excel', query_string={'product_id': product['id']})
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data))
    expected = {'Assembly': 24, 'Nested': 38, 'Part': 124, 'Sibling': 124}
    for sheet in workbook:
        for row in sheet.iter_rows(min_row=2, values_only=True):
            assert row[4] == expected[row[0].strip()]
    client.patch('/api/node/' + product['id'], json={'order_count': 0})
    response = client.get('/api/export/excel', query_string={'product_id': product['id']})
    assert 'OrderQty0' in response.headers['Content-Disposition']
    for sheet in load_workbook(BytesIO(response.data)):
        assert all(row[4] == 0 for row in sheet.iter_rows(min_row=2, values_only=True))


def test_pdf_labels_use_shared_calculation(client, monkeypatch):
    product, assembly, nested, part, sibling = tree(client)
    captured = []
    def pipe(graph):
        captured.append(graph.source)
        return b'%PDF-1.4 mock'
    monkeypatch.setattr(backend.graphviz.Digraph, 'pipe', pipe)
    response = client.get('/api/export/schematic', query_string={'product_id': product['id']})
    assert response.status_code == 200
    assert 'نیاز کل: 124 | کسری: 118' in captured[0]
    assert 'نیاز کل: 38 | کسری: 31' in captured[0]
    client.patch('/api/node/' + product['id'], json={'order_count': 0})
    client.get('/api/export/schematic', query_string={'product_id': product['id']})
    assert 'نیاز کل: 124' not in captured[1]


def test_v2_inventory_writes_and_schedule_preservation(client):
    user = User(username='isolated-admin', email='isolated@example.com', role='admin', password_hash='test-only')
    db.session.add(user)
    db.session.commit()
    with client.session_transaction() as session:
        session['_user_id'] = str(user.id)
        session['_fresh'] = True
    response = client.post('/api/v2/products', json={'name': 'V2', 'quantity': '2', 'order_count': 10})
    assert response.status_code == 201
    product = response.json['data']
    response = client.post('/api/v2/assemblies', json={
        'name': 'V2 Assembly', 'product_id': product['id'], 'quantity': 3, 'required_quantity': 4})
    assert response.status_code == 201
    assembly = response.json['data']
    response = client.post('/api/v2/assemblies', json={
        'name': 'V2 Nested', 'product_id': product['id'], 'parent_id': assembly['id'],
        'quantity': 1, 'required_quantity': 2})
    assert response.status_code == 201
    nested = response.json['data']
    response = client.post('/api/v2/parts', json={
        'name': 'V2 Part', 'assembly_id': nested['id'], 'quantity': 5, 'required_quantity': 7})
    assert response.status_code == 201
    part = response.json['data']
    assert client.put('/api/v2/products/' + str(product['id']), json={'quantity': -1}).status_code == 400
    assert client.patch('/api/v2/assemblies/' + str(assembly['id']), json={'required_quantity': 1.5}).status_code == 400
    assert client.put('/api/v2/parts/' + str(part['id']), json={'quantity': True}).status_code == 400
    response = client.patch('/api/v2/assemblies/' + str(assembly['id']), json={'quantity': 4, 'required_quantity': 3})
    assert response.json['data']['quantity'] == 4
    before = client.get('/api/data').json
    response = client.post('/api/v2/schedules', json={'product_id': product['id'], 'quantity': 9})
    assert response.status_code == 201
    assert response.json['data']['quantity'] == 9
    assert client.get('/api/data').json == before
    orders = client.get('/api/v2/work-orders').json['data']
    # B10: nested multipliers count, so 3 (assembly) × 2 (nested) × 7 (part) × 9 units.
    assert len(orders) == 1 and orders[0]['quantity'] == 378
    assert client.get('/api/schedules').json['data'][0]['quantity'] == 9


def test_email_values_and_zero_shortage_skip(client, monkeypatch):
    import email_notif
    product, assembly, nested, part, sibling = tree(client)
    values = []
    sent = []
    monkeypatch.setattr(email_notif, 'render_template', lambda template, **kwargs: values.append(kwargs) or 'email')
    monkeypatch.setattr(email_notif.mail, 'send', lambda message: sent.append(message))
    model = db.session.get(Part, int(part['id'][1:]))
    assert email_notif.send_low_stock_email(model, recipient='override@example.com')
    assert values[0]['required'] == 124 and values[0]['shortage'] == 118
    assert sent[0].recipients == ['override@example.com']
    client.patch('/api/node/' + product['id'], json={'order_count': 0})
    assert not email_notif.send_low_stock_email(model)
    email_notif.check_and_notify_low_stock()
    assert len(sent) == 1
