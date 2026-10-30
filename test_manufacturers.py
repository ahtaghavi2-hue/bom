"""Manufacturer regression tests; never open the configured production database."""
from io import BytesIO
from pathlib import Path
import os
import sqlite3
import subprocess
import sys
from unittest.mock import patch

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, inspect, text
from config import Config

with patch.object(Config, 'SQLALCHEMY_DATABASE_URI', 'sqlite:///:memory:'):
    import app as backend

from extensions import db
from models import (Manufacturer, ManufacturerEmail, ManufacturerPhone, ManufacturerSocial,
                    Product, Assembly, Part, Stage, StageDetail, StageManufacturer,
                    Category, Tag, User)
from manufacturer_migration import ensure_manufacturer_schema


@pytest.fixture
def client():
    app = backend.app
    config = dict(app.config)
    keys = ('bom_inventory_schema_ready', 'manufacturer_schema_ready')
    flags = {key: app.extensions[key] for key in keys if key in app.extensions}
    try:
        app.config.update(TESTING=True, LOGIN_DISABLED=False)
        with app.app_context():
            assert db.engine.url.database == ':memory:'
            try:
                db.create_all()
                with db.engine.begin() as conn:
                    conn.execute(text('DROP TABLE IF EXISTS manufacturer_schema_state'))
                for key in keys:
                    app.extensions.pop(key, None)
                for role in ('admin', 'engineer', 'viewer'):
                    user = User(username=role, email=role + '@example.com', role=role)
                    user.set_password('test-password')
                    db.session.add(user)
                db.session.commit()
                client = app.test_client()
                login(client, 'admin')
                yield client
            finally:
                db.session.remove()
                db.drop_all()
                with db.engine.begin() as conn:
                    conn.execute(text('DROP TABLE IF EXISTS manufacturer_schema_state'))
    finally:
        app.config.clear()
        app.config.update(config)
        for key in keys:
            app.extensions.pop(key, None)
        app.extensions.update(flags)


def login(client, role):
    from flask import g
    g.pop('_login_user', None)
    user = User.query.filter_by(username=role).one()
    with client.session_transaction() as session:
        session.clear()
        session['_user_id'] = str(user.id)
        session['_fresh'] = True


def post(client, path, payload):
    response = client.post('/api/v2/' + path, json=payload)
    assert response.status_code == 201, response.get_data(as_text=True)
    return response.json['data']


def manufacturer(client, **values):
    return post(client, 'manufacturers', {'name': 'Maker', **values})


def part_tree():
    product = Product(name='=Product')
    assembly = Assembly(name='Assembly', product=product)
    part = Part(name='@Part', assembly=assembly)
    db.session.add(product)
    db.session.commit()
    return part


def test_crud_contacts_omission_and_legacy_dedup(client):
    data = manufacturer(client, phone='00123', address='old address', notes='old notes',
        emails=[{'email': 'a@example.com'}], phones=[{'phone': '00123'}, {'phone': '00456'}, {'phone': '00456'}],
        socials=[{'platform': 'web', 'handle': 'example.com'}])
    mid = data['id']
    assert [item['phone'] for item in data['phones']] == ['00123', '00456']
    response = client.put(f'/api/v2/manufacturers/{mid}', json={'name': 'Renamed', 'is_active': False})
    assert response.status_code == 200
    updated = response.json['data']
    for key in ('phone', 'phones', 'emails', 'socials', 'address', 'notes'):
        assert updated[key] == data[key]
    response = client.put(f'/api/v2/manufacturers/{mid}', json={'phones': []})
    assert response.json['data']['phones'] == [{'id': None, 'phone': '00123'}]
    assert db.session.get(Manufacturer, mid).phone == '00123'
    response = client.put(f'/api/v2/manufacturers/{mid}', json={
        'emails': [{'email': 'replacement@example.com'}], 'socials': []})
    assert response.status_code == 200
    assert [item['email'] for item in response.json['data']['emails']] == ['replacement@example.com']
    assert response.json['data']['socials'] == []
    assert client.get(f'/api/v2/manufacturers/{mid}').json['data']['usage_stages'] == []
    assert client.delete(f'/api/v2/manufacturers/{mid}').status_code == 200
    assert client.get(f'/api/v2/manufacturers/{mid}').status_code == 404


@pytest.mark.parametrize('bad', [
    {'emails': None}, {'phones': ['x']}, {'socials': [{'platform': 'x'}]},
    {'emails': [{'email': 1}]}, {'phones': [{'phone': ''}]}, {'is_active': 'false'},
    {'tag_ids': [999999]}, {'tag_ids': [True]}, {'tag_ids': None}, {'name': ''},
    {'phone': 123}, {'address': []}, {'notes': {}}, {'phones': [{'phone': 'x' * 51}]},
])
def test_malformed_manufacturer_atomic(client, bad):
    original = manufacturer(client, emails=[{'email': 'old@example.com'}])
    payload = {'name': 'MUTATED', 'emails': [{'email': 'new@example.com'}], **bad}
    assert client.put(f"/api/v2/manufacturers/{original['id']}", json=payload).status_code == 400
    assert client.get(f"/api/v2/manufacturers/{original['id']}").json['data']['name'] == original['name']
    assert client.get(f"/api/v2/manufacturers/{original['id']}").json['data']['emails'] == original['emails']
    assert client.post('/api/v2/manufacturers', json=payload).status_code == 400
    assert Manufacturer.query.count() == 1


def test_invalid_json(client):
    for payload in ([], None, 'text'):
        assert client.post('/api/v2/manufacturers', json=payload).status_code == 400
    assert client.post('/api/v2/manufacturer-tags', json=[]).status_code == 400


def test_definitions_filters_and_guards(client):
    first = post(client, 'manufacturer-categories', {'name': 'Custom process'})
    second = post(client, 'manufacturer-categories', {'name': 'Custom quality'})
    tag1 = post(client, 'manufacturer-tags', {'name': 'Same label', 'category_id': first['id']})
    tag2 = post(client, 'manufacturer-tags', {'name': 'Same label', 'category_id': second['id']})
    assert client.post('/api/v2/manufacturer-tags', json={'name': 'Same label', 'category_id': first['id']}).status_code == 400
    both = manufacturer(client, name='Both', tag_ids=[tag1['id'], tag2['id']], phone='00999')
    only = manufacturer(client, name='Only', tag_ids=[tag1['id']], is_active=False)
    def ids(url):
        return [item['id'] for item in client.get('/api/v2/manufacturers' + url).json['data']]
    assert set(ids('')) == {both['id'], only['id']}
    assert ids('?active=true') == [both['id']]
    assert ids('?active=false') == [only['id']]
    assert ids(f"?tag_id={tag1['id']}&tag_id={tag2['id']}") == [both['id']]
    assert ids(f"?category_id={second['id']}") == [both['id']]
    assert ids('?q=00999') == [both['id']]
    assert ids('?q=Same%20label&active=false') == [only['id']]
    for query in ('?active=bad', '?tag_id=x', '?category_id=0'):
        assert client.get('/api/v2/manufacturers' + query).status_code == 400
        assert client.get('/api/v2/manufacturers/export/excel' + query).status_code == 400
    assert client.delete(f"/api/v2/manufacturer-tags/{tag1['id']}").status_code == 409
    assert client.delete(f"/api/v2/manufacturer-categories/{first['id']}").status_code == 409
    assert client.put(f"/api/v2/manufacturer-tags/{tag2['id']}", json={'name': 'Changed'}).status_code == 200
    assert client.put(f"/api/v2/manufacturer-categories/{second['id']}", json={'name': 'Changed category'}).status_code == 200
    data = client.get(f"/api/v2/manufacturers/{both['id']}").json['data']
    assert data['tags'][1]['category_name'] == 'Changed category'
    assert client.put(f"/api/v2/manufacturers/{both['id']}", json={'tag_ids': []}).status_code == 200
    assert client.delete(f"/api/v2/manufacturer-tags/{tag2['id']}").status_code == 200
    assert client.delete(f"/api/v2/manufacturer-categories/{second['id']}").status_code == 200


def test_ordered_unlimited_independent_stages_and_classic_save(client):
    mids = [manufacturer(client, name=f'Maker {i}')['id'] for i in range(15)]
    part = part_tree()
    stages = [post(client, 'stages', {'part_id': part.id, 'name': name, 'manufacturer_ids': values})
              for name, values in [('Cut', mids), ('Drill', mids[::-1])]]
    detail = StageDetail(stage_id=stages[0]['id'], description='Keep detail')
    db.session.add(detail)
    db.session.commit()
    for stage, values in zip(stages, (mids, mids[::-1])):
        data = client.get(f"/api/v2/stages/{stage['id']}/manufacturers").json['data']
        assert data['manufacturer_ids'] == values
        assert [item['priority'] for item in data['manufacturers']] == list(range(1, 16))
        assert 'effective_manufacturer' not in data
    sid = stages[0]['id']
    assert client.put(f'/api/v2/stages/{sid}/manufacturers', json={'manufacturer_ids': mids[::-1]}).status_code == 200
    assert client.put(f"/api/v2/manufacturers/{mids[0]}", json={'is_active': False}).status_code == 200
    assert client.delete(f"/api/v2/manufacturers/{mids[0]}").status_code == 409
    data = client.get(f'/api/v2/stages/{sid}/manufacturers').json['data']
    assert data['manufacturer_ids'] == mids[::-1]
    assert data['manufacturers'][-1]['is_active'] is False
    response = client.put(f'/api/node/r{part.id}', json={'stages': [
        {'id': stages[1]['id']}, {'id': sid, 'manufacturer_ids': mids}]})
    assert response.status_code == 200
    assert [stage['id'] for stage in response.json['node']['stages']] == [stages[1]['id'], sid]
    assert response.json['node']['stages'][1]['manufacturer_ids'] == mids
    assert db.session.get(StageDetail, detail.id).description == 'Keep detail'
    assert part.manufacturers.count() == 0
    classic_new = client.post('/api/node', json={'type': 'part', 'parent': f'a{part.assembly_id}',
        'name': 'New', 'stages': [{'name': 'Fresh', 'manufacturer_ids': mids}]})
    assert classic_new.status_code == 200
    assert classic_new.json['node']['stages'][0]['manufacturer_ids'] == mids
    assert client.get(f"/api/v2/manufacturers/{mids[0]}").json['data']['usage_stages'][0]['product_name'] == '=Product'
    assert client.delete(f'/api/v2/stages/{sid}').status_code == 200
    assert StageManufacturer.query.filter_by(stage_id=sid).count() == 0
    assert db.session.get(StageDetail, detail.id) is None


@pytest.mark.parametrize('bad', [[1, 1], [999999], [True], ['1'], None, {}])
def test_stage_nested_atomic_rejection(client, bad):
    mid = manufacturer(client)['id']
    part = part_tree()
    stages = [post(client, 'stages', {'part_id': part.id, 'name': name, 'manufacturer_ids': [mid]})
              for name in ('First', 'Second')]
    response = client.put(f'/api/node/r{part.id}', json={'name': 'MUTATED', 'stages': [
        {'id': stages[0]['id'], 'name': 'MUTATED', 'manufacturer_ids': []},
        {'id': stages[1]['id'], 'manufacturer_ids': bad}]})
    assert response.status_code == 400
    assert db.session.get(Part, part.id).name == '@Part'
    assert db.session.get(Stage, stages[0]['id']).name == 'First'
    assert client.get(f"/api/v2/stages/{stages[0]['id']}/manufacturers").json['data']['manufacturer_ids'] == [mid]
    assert client.post('/api/v2/stages', json={'part_id': part.id, 'name': 'Bad', 'manufacturer_ids': bad}).status_code == 400
    assert client.put(f"/api/v2/stages/{stages[0]['id']}", json={'name': 'Bad', 'manufacturer_ids': bad}).status_code == 400
    assert Stage.query.count() == 2


def test_part_reference_deletion_guard_no_stage_inference(client):
    mid = manufacturer(client)['id']
    part = part_tree()
    assert client.post(f'/api/v2/parts/{part.id}/manufacturers', json={'manufacturer_id': mid}).status_code == 201
    stage = post(client, 'stages', {'part_id': part.id, 'name': 'No inferred link'})
    assert stage['manufacturer_ids'] == []
    assert client.delete(f'/api/v2/manufacturers/{mid}').status_code == 409
    assert client.delete(f'/api/v2/parts/{part.id}/manufacturers/{mid}').status_code == 200
    assert client.delete(f'/api/v2/manufacturers/{mid}').status_code == 200


@pytest.mark.parametrize('role', ['admin', 'engineer', 'viewer'])
def test_auth_roles(client, role):
    mid = manufacturer(client)['id']
    part = part_tree()
    stage = post(client, 'stages', {'part_id': part.id, 'name': 'Stage'})
    login(client, role)
    assert client.get('/api/v2/manufacturers').status_code == 200
    assert client.get('/api/v2/manufacturer-categories').status_code == 200
    expected = 403 if role == 'viewer' else 201
    assert client.post('/api/v2/manufacturers', json={'name': role}).status_code == expected
    assert client.post('/api/v2/manufacturer-categories', json={'name': role}).status_code == expected
    assert client.put(f'/api/v2/manufacturers/{mid}', json={'is_active': False}).status_code == (403 if role == 'viewer' else 200)
    assert client.put(f"/api/v2/stages/{stage['id']}/manufacturers", json={'manufacturer_ids': [mid]}).status_code == (403 if role == 'viewer' else 200)
    category = client.get('/api/v2/manufacturer-categories').json['data'][0]
    assert client.post('/api/v2/manufacturer-tags', json={
        'name': role, 'category_id': category['id']}).status_code == expected
    assert client.put(f"/api/v2/manufacturer-categories/{category['id']}", json={
        'name': 'Renamed by ' + role}).status_code == (403 if role == 'viewer' else 200)
    assert client.delete(f'/api/v2/manufacturers/{mid}').status_code == (403 if role == 'viewer' else 409)
    assert client.delete(f"/api/v2/manufacturer-categories/{category['id']}").status_code == (403 if role == 'viewer' else 409)
    with client.session_transaction() as session:
        session.clear()
    from flask import g
    g.pop('_login_user', None)
    for path in ('/api/v2/manufacturers', '/api/v2/manufacturer-tags', '/manufacturers', '/api/v2/manufacturers/export/excel'):
        assert client.get(path).status_code == 302


@pytest.mark.parametrize('assignment', ['replace', 'clear', 'reorder', 'new', 'legacy', 'invalid'])
def test_viewer_classic_stage_changes_forbidden_before_any_mutation(client, assignment):
    from werkzeug.exceptions import Forbidden

    mids = [manufacturer(client, name=f'Maker {i}')['id'] for i in range(3)]
    part = part_tree()
    stages = [post(client, 'stages', {'part_id': part.id, 'name': name, 'manufacturer_ids': mids[:2]})
              for name in ('First', 'Second')]
    before = client.get('/api/data').json['nodes'][f'r{part.id}']
    login(client, 'viewer')
    entries = [{'id': stages[0]['id'], 'name': 'MUTATED', 'manufacturer_ids': mids[:2]},
               {'id': stages[1]['id'], 'manufacturer_ids': mids[:2]}]
    entries[1]['manufacturer_ids'] = {
        'replace': [mids[2]], 'clear': [], 'reorder': mids[:2][::-1],
        'new': [mids[2]], 'legacy': [mids[2]], 'invalid': [999999],
    }[assignment]
    if assignment == 'new':
        entries[1].pop('id')
    elif assignment == 'legacy':
        for entry in entries:
            entry.pop('id')
    payload = {'name': 'MUTATED', 'quantity': 99, 'stages': entries}
    # Verify preflight itself, rather than relying only on endpoint rollback.
    with backend.app.test_request_context():
        from flask import g
        g._login_user = User.query.filter_by(username='viewer').one()
        with pytest.raises(Forbidden):
            backend._apply_node_fields(part, 'part', payload)
        assert part.name == before['name'] and part.quantity == before['quantity']
        assert not db.session.dirty and not db.session.new and not db.session.deleted
    response = client.put(f'/api/node/r{part.id}', json=payload)
    assert response.status_code == 403
    assert 'error' in response.json
    assert client.get('/api/data').json['nodes'][f'r{part.id}'] == before
    assert Stage.query.count() == 2


@pytest.mark.parametrize('role', ['admin', 'engineer', 'viewer'])
def test_classic_new_stage_assignment_role_policy(client, role):
    mid = manufacturer(client)['id']
    part = part_tree()
    login(client, role)
    response = client.post('/api/node', json={'type': 'part', 'parent': f'a{part.assembly_id}',
        'name': 'New', 'stages': [{'name': 'First'}, {'name': 'Assigned', 'manufacturer_ids': [mid]}]})
    assert response.status_code == (403 if role == 'viewer' else 200)
    assert Part.query.count() == (1 if role == 'viewer' else 2)
    assert Stage.query.count() == (0 if role == 'viewer' else 2)
    if role != 'viewer':
        assert response.json['node']['stages'][1]['manufacturer_ids'] == [mid]


@pytest.mark.parametrize('legacy', [False, True])
def test_viewer_classic_edits_denied_engineer_legacy_edits_allowed(client, legacy):
    mids = [manufacturer(client, name=f'Maker {i}')['id'] for i in range(2)]
    part = part_tree()
    stage = post(client, 'stages', {'part_id': part.id, 'name': 'Stage', 'manufacturer_ids': mids})
    # B2: classic mutations are admin/engineer only, so a viewer is refused.
    login(client, 'viewer')
    serialized = client.get('/api/data').json['nodes'][f'r{part.id}']
    serialized['notes'] = 'Viewer legacy edit'
    if legacy:
        serialized['stages'][0].pop('id')
    assert client.put(f'/api/node/r{part.id}', json=serialized).status_code == 403
    assert client.patch(f'/api/node/r{part.id}', json={'stages': [{'id': stage['id'], 'name': 'Renamed'}]}).status_code == 403
    assert client.post('/api/node', json={'type': 'part', 'parent': f'a{part.assembly_id}',
        'stages': [{'name': 'Unassigned', 'manufacturer_ids': []}]}).status_code == 403
    # The same legacy payloads still work for an engineer and keep associations.
    login(client, 'engineer')
    response = client.put(f'/api/node/r{part.id}', json=serialized)
    assert response.status_code == 200
    assert response.json['node']['notes'] == 'Viewer legacy edit'
    assert response.json['node']['stages'][0]['manufacturer_ids'] == mids
    response = client.patch(f'/api/node/r{part.id}', json={'stages': [{'id': stage['id'], 'name': 'Renamed'}]})
    assert response.status_code == 200
    assert response.json['node']['stages'][0]['manufacturer_ids'] == mids
    assert client.post('/api/node', json={'type': 'part', 'parent': f'a{part.assembly_id}',
        'stages': [{'name': 'Unassigned', 'manufacturer_ids': []}]}).status_code == 200


def test_excel_contents_filters_formula_safety(client):
    category = post(client, 'manufacturer-categories', {'name': '=Category'})
    tag = post(client, 'manufacturer-tags', {'name': '=Tag', 'category_id': category['id']})
    data = manufacturer(client, name='=Maker', phone='001234', notes='=HYPERLINK("bad")',
        address='@Address', emails=[{'email': '=email'}], socials=[{'platform': '=web', 'handle': '+handle'}], tag_ids=[tag['id']])
    manufacturer(client, name='Excluded', is_active=False)
    part = part_tree()
    post(client, 'stages', {'part_id': part.id, 'name': '=Stage', 'manufacturer_ids': [data['id']]})
    for path in (f"/api/v2/manufacturers/{data['id']}/export/excel",
                 f"/api/v2/manufacturers/export/excel?active=true&tag_id={tag['id']}&q=Maker"):
        response = client.get(path)
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.data), data_only=False)
        assert workbook.sheetnames == ['سازندگان', 'اطلاعات تماس', 'برچسب‌ها', 'مراحل استفاده']
        assert workbook['سازندگان'].max_row == 2
        assert workbook['سازندگان']['B2'].value == '=Maker'
        assert workbook['اطلاعات تماس']['D2'].value == '001234'
        assert workbook['برچسب‌ها']['D2'].value == '=Tag'
        assert [workbook['مراحل استفاده'].cell(2, col).value for col in range(3, 7)] == ['=Product', '@Part', '=Stage', 1]
        for sheet in workbook:
            assert sheet.sheet_view.rightToLeft and sheet.freeze_panes == 'A2'
            assert sheet.auto_filter.ref == sheet.dimensions
            for row in sheet:
                for cell in row:
                    assert cell.font.name == 'Arial'
                    assert cell.data_type != 'f'
                    if isinstance(cell.value, str):
                        assert cell.data_type == 's'


@pytest.mark.parametrize('suite', ['inventory', 'manufacturers'])
@pytest.mark.parametrize('fail', [False, True])
def test_fixture_restores_config_flags_and_reseeds_isolated_database(suite, fail):
    import test_bom_inventory

    app = backend.app
    original_config = dict(app.config)
    keys = ('bom_inventory_schema_ready', 'manufacturer_schema_ready')
    original_flags = {key: app.extensions[key] for key in keys if key in app.extensions}
    fixture = test_bom_inventory.client if suite == 'inventory' else client
    try:
        app.config['LOGIN_DISABLED'] = False
        app.config['TESTING'] = False
        app.extensions['bom_inventory_schema_ready'] = 'preserve-inventory'
        app.extensions['manufacturer_schema_ready'] = 'preserve-manufacturers'
        expected_config = dict(app.config)
        for attempt in range(2):
            generator = fixture.__wrapped__()
            isolated_client = next(generator)
            try:
                assert isolated_client.get('/api/v2/manufacturer-categories').status_code == 200
                assert Category.query.count() == 4
                with db.engine.connect() as conn:
                    assert 'manufacturer_schema_state' in inspect(conn).get_table_names()
            finally:
                if fail:
                    with pytest.raises(RuntimeError, match='fixture failure'):
                        generator.throw(RuntimeError('fixture failure'))
                else:
                    with pytest.raises(StopIteration):
                        next(generator)
            assert dict(app.config) == expected_config
            assert app.extensions['bom_inventory_schema_ready'] == 'preserve-inventory'
            assert app.extensions['manufacturer_schema_ready'] == 'preserve-manufacturers'
            with app.app_context():
                assert db.engine.url.database == ':memory:'
                assert 'manufacturer_schema_state' not in inspect(db.engine).get_table_names()
    finally:
        app.config.clear()
        app.config.update(original_config)
        for key in keys:
            app.extensions.pop(key, None)
        app.extensions.update(original_flags)


def test_migration_backup_idempotence_and_deleted_seeds(tmp_path):
    path = tmp_path / 'legacy.sqlite'
    with sqlite3.connect(path) as conn:
        conn.executescript('CREATE TABLE manufacturers (id INTEGER PRIMARY KEY, name TEXT, phone TEXT);'
                          "INSERT INTO manufacturers VALUES (1, 'Old', '00123');"
                          'CREATE TABLE untouched (value TEXT);'
                          "INSERT INTO untouched VALUES ('keep');")
    engine = create_engine('sqlite:///' + path.as_posix())
    backup = ensure_manufacturer_schema(engine, tmp_path / 'backups')
    assert Path(backup).exists()
    with sqlite3.connect(backup) as conn:
        assert 'is_active' not in [row[1] for row in conn.execute('PRAGMA table_info(manufacturers)')]
        assert conn.execute('SELECT * FROM untouched').fetchall() == [('keep',)]
    with engine.begin() as conn:
        assert conn.execute(text('SELECT name,phone,is_active FROM manufacturers')).one() == ('Old', '00123', 1)
        assert conn.execute(text('SELECT count(*) FROM manufacturer_categories')).scalar() == 4
        conn.execute(text('DELETE FROM manufacturer_tags_definitions'))
        conn.execute(text('DELETE FROM manufacturer_categories'))
    assert ensure_manufacturer_schema(engine, tmp_path / 'backups') is None
    with engine.connect() as conn:
        assert conn.execute(text('SELECT count(*) FROM manufacturer_categories')).scalar() == 0
        assert conn.execute(text('SELECT * FROM untouched')).one() == ('keep',)
    engine.dispose()


def test_backup_failure_prevents_schema_changes(tmp_path):
    path = tmp_path / 'failure.sqlite'
    with sqlite3.connect(path) as conn:
        conn.execute('CREATE TABLE manufacturers (id INTEGER PRIMARY KEY, name TEXT)')
    engine = create_engine('sqlite:///' + path.as_posix())
    with patch('manufacturer_migration.sqlite3.connect', side_effect=OSError('Backup denied')):
        with pytest.raises(OSError, match='Backup denied'):
            ensure_manufacturer_schema(engine, tmp_path / 'backups')
    assert inspect(engine).get_table_names() == ['manufacturers']
    assert 'is_active' not in {column['name'] for column in inspect(engine).get_columns('manufacturers')}
    engine.dispose()


def test_existing_definitions_never_seeded(tmp_path):
    engine = create_engine('sqlite:///:memory:')
    Category.__table__.create(engine)
    with engine.begin() as conn:
        conn.execute(Category.__table__.insert().values(name='User category'))
    ensure_manufacturer_schema(engine)
    with engine.connect() as conn:
        assert conn.execute(text('SELECT name FROM manufacturer_categories')).fetchall() == [('User category',)]
        assert conn.execute(text('SELECT count(*) FROM manufacturer_tags_definitions')).scalar() == 0
    engine.dispose()


def test_actual_startup_and_first_request_migrate_before_user_load(tmp_path):
    # Exercise the real app import and init_db in fresh processes, with URI set
    # before import. The existing production DB is never attached or queried.
    path = tmp_path / 'startup.sqlite'
    env = dict(os.environ, DATABASE_URL='sqlite:///' + path.as_posix(), PYTHONIOENCODING='utf-8')
    root = Path(__file__).resolve().parent
    code = '''
from pathlib import Path
import app as backend
from extensions import db
from sqlalchemy import text
from unittest.mock import patch
with backend.app.app_context():
    assert Path(db.engine.url.database).resolve() == Path(__import__('os').environ['DATABASE_URL'][10:]).resolve()
    from manufacturer_migration import ensure_manufacturer_schema
    def migrate(engine):
        return ensure_manufacturer_schema(engine, Path(engine.url.database).parent / 'backups')
    with patch.object(backend, 'ensure_manufacturer_schema', migrate):
        backend.init_db()
        with db.engine.begin() as conn:
            conn.execute(text('ALTER TABLE manufacturers DROP COLUMN is_active'))
        backend.app.extensions.pop('manufacturer_schema_ready', None)
        client = backend.app.test_client()
        with client.session_transaction() as session:
            session['_user_id'] = '1'
        assert client.get('/api/v2/manufacturers').status_code == 200
    assert 'manufacturers_page' in backend.app.view_functions
    with patch.object(backend, 'render_template', return_value='dashboard') as render:
        assert client.get('/manufacturers').data == b'dashboard'
        render.assert_called_once_with('manufacturers.html')
    backend.init_db()
print('REAL_STARTUP_OK')
'''
    result = subprocess.run([sys.executable, '-c', code], cwd=root, env=env, capture_output=True, text=True)
    assert result.returncode == 0, result.stdout + result.stderr
    assert 'REAL_STARTUP_OK' in result.stdout
    assert list((tmp_path / 'backups').glob('*.sqlite'))
